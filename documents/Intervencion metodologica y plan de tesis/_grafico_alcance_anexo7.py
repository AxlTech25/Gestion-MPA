# -*- coding: utf-8 -*-
"""Gráfico de alcance (incluido / fuera) e inserción en el apartado 1.3."""
from pathlib import Path
from textwrap import wrap

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent
IMG = BASE / "_figuras_anexo7" / "fig_a7_alcance.png"
DOC = BASE / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

PARRAFO = (
    "El alcance de esta intervención se delimita a la construcción documentada de "
    "Sigemad MPA V2 con el Prompt-Centered SDLC v1.2, desde la elicitación de "
    "requisitos hasta la puesta en producción. El Gráfico 1 distingue lo incluido "
    "de lo excluido. Los macroprocesos institucionales y los diagramas de proceso "
    "se desarrollan en los apartados 3.4 y 3.5."
)
CAPTION = "Gráfico 1. Alcance de la intervención: aspectos incluidos y fuera de alcance."

INCLUIDO = [
    "Cinco fases del SDLC (R, D, I, T, M)",
    "Agente acotado con revisión humana (N2)",
    "Repositorio de prompts vigentes",
    "Matriz de doble entrada función–prompt–commit",
    "Puesta en producción institucional, con ML opcional (ADR-002)",
]
FUERA = [
    "Entrenamiento de un modelo fundacional propio",
    "Automatización de niveles N4–N5",
    "Portal ciudadano, bienes no informáticos y SIGA/SIAF",
    "Contrastación de hipótesis y objetivos del Capítulo I",
]


def _font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK


def style_body(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0.75)
    for run in p.runs:
        _font(run, 12)


def insert_after(element, doc):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    return Paragraph(new_p, doc)


def draw_graphic():
    plt.rcParams["font.family"] = "Times New Roman"
    fig, ax = plt.subplots(figsize=(10.4, 5.6), dpi=170)
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 10.4)
    ax.set_ylim(0, 5.6)
    ax.axis("off")

    ax.text(
        5.2, 5.28, "Alcance de la intervención",
        ha="center", va="center", fontsize=14, fontweight="bold", color="#1F3A5F",
    )

    cols = [
        (0.28, "#1F3A5F", "#D6E4F0", "Incluido", INCLUIDO),
        (5.32, "#6B2D3C", "#F3E8EA", "Fuera de alcance", FUERA),
    ]
    col_w, col_h = 4.80, 4.55
    y0 = 0.22

    for x, header, fill, title, items in cols:
        ax.add_patch(
            FancyBboxPatch(
                (x, y0), col_w, col_h,
                boxstyle="round,pad=0.02,rounding_size=0.08",
                linewidth=1.4, edgecolor=header, facecolor=fill,
            )
        )
        ax.add_patch(Rectangle((x, y0 + col_h - 0.72), col_w, 0.72, facecolor=header, edgecolor=header))
        ax.text(
            x + col_w / 2, y0 + col_h - 0.36, title,
            ha="center", va="center", fontsize=13, fontweight="bold", color="white",
        )
        yy = y0 + col_h - 1.02
        for item in items:
            wrapped = "\n".join(wrap(item, width=38))
            ax.text(x + 0.22, yy, "•", ha="left", va="top", fontsize=12, color=header)
            ax.text(
                x + 0.48, yy, wrapped,
                ha="left", va="top", fontsize=10.0, color="#1A1A1A",
                linespacing=1.25,
            )
            yy -= 0.68 if "\n" not in wrapped else 0.88

    IMG.parent.mkdir(exist_ok=True)
    fig.savefig(IMG, bbox_inches="tight", facecolor="white", pad_inches=0.12)
    plt.close()


def replace_paragraph_text(p, text, size=12, italic=False, center=False):
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size, italic=italic)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size, italic=italic)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
    else:
        style_body(p)


def already_has_grafico(doc):
    return any("Gráfico 1. Alcance de la intervención" in p.text for p in doc.paragraphs)


def insert_in_docx():
    doc = Document(str(DOC))
    if already_has_grafico(doc):
        for p in doc.paragraphs:
            if p.text.strip().startswith("El alcance de"):
                replace_paragraph_text(p, PARRAFO)
                break
        doc.save(str(DOC))
        print("Ya estaba insertado; se actualizo el parrafo.")
        return

    heading = None
    body = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "1.3 Alcance":
            heading = p
            for q in doc.paragraphs[i + 1 : i + 6]:
                if q.text.strip():
                    body = q
                    break
            break
    if heading is None:
        raise RuntimeError("No se encontro 1.3 Alcance")

    if body is not None and body.text.strip().startswith("El alcance"):
        replace_paragraph_text(body, PARRAFO)
        anchor = body._element
    else:
        p = insert_after(heading._element, doc)
        run = p.add_run(PARRAFO)
        _font(run)
        style_body(p)
        anchor = p._element

    pic_p = insert_after(anchor, doc)
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(6)
    pic_p.paragraph_format.space_after = Pt(2)
    pic_p.paragraph_format.first_line_indent = Cm(0)
    pic_p.add_run().add_picture(str(IMG), width=Cm(15.5))

    cap = insert_after(pic_p._element, doc)
    run = cap.add_run(CAPTION)
    _font(run, size=11, italic=True)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.0

    doc.save(str(DOC))
    print("OK", DOC)


if __name__ == "__main__":
    draw_graphic()
    insert_in_docx()
    print("PNG", IMG)

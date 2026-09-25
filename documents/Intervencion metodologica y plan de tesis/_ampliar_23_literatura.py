# -*- coding: utf-8 -*-
"""Amplía el apartado 2.3 a cinco fuentes y actualiza la Tabla 4 y las referencias."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent
DOC = BASE / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)


def _font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_body(p, size=12, first_indent=0.75, after=8, justify=True, center=False, bold=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    )
    pf = p.paragraph_format
    pf.space_before = Pt(6 if bold and not first_indent else 0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0 if center or bold else first_indent)


def set_runs_text(p, text, size=12, bold=False, italic=False, justify=True, first_indent=0.75, after=8, center=False):
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size, bold=bold, italic=italic)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size, bold=bold, italic=italic)
    style_body(p, size=size, first_indent=first_indent, after=after, justify=justify, center=center, bold=bold)


def insert_after(element, doc):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    return Paragraph(new_p, doc)


def insert_para(el, text, doc, size=12, bold=False, first_indent=0.75, after=8, justify=True):
    p = insert_after(el, doc)
    run = p.add_run(text)
    _font(run, size=size, bold=bold)
    style_body(
        p,
        size=size,
        first_indent=0 if bold else first_indent,
        after=after,
        justify=False if bold else justify,
        bold=bold,
    )
    return p


def find_body_heading(doc, text):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    return matches[-1] if matches else None


def has_text(doc, needle):
    return any(needle in p.text for p in doc.paragraphs)


def set_cell_text(cell, text, size=8, bold=False):
    p = cell.paragraphs[0]
    for extra in list(cell.paragraphs)[1:]:
        extra._element.getparent().remove(extra._element)
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size, bold=bold)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size, bold=bold)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


INTRO = (
    "El apartado 2.1 situó el problema (ambigüedad del pedido al modelo, código inseguro y "
    "desajuste de Scrum para un solo desarrollador con agente). Este apartado posiciona la "
    "intervención frente a cinco fuentes de la literatura de inteligencia artificial en el ciclo "
    "de vida del software. Se eligieron porque cubren, sin repetirse, las cinco pretensiones del "
    "Prompt-Centered SDLC: conservar fases estables, integrar IA con gobernanza, abandonar el "
    "sprint como unidad de planificación, tratar el prompt como artefacto de calidad y evaluar "
    "las salidas generadas antes de integrarlas. La Tabla 4 resume, para cada fuente, qué aporta, "
    "qué no cubre y qué brecha cierra este caso."
)

SECTIONS = [
    (
        "2.3.1 Yas, Alazzawi y Rahmatullah (2023): el esqueleto del SDLC",
        [
            (
                "Yas, Alazzawi y Rahmatullah [22] revisan metodologías del ciclo de vida y las "
                "agrupan en enfoques tradicionales (cascada, iterativo, espiral, V-Model) y "
                "ágiles (XP, Scrum, FDD, Kanban). El hallazgo que se retiene no es la superioridad "
                "de un modelo, sino dos hechos: no existe metodología universal, y todo desarrollo "
                "sigue fases reconocibles de análisis de requisitos, diseño, implementación, "
                "pruebas y mantenimiento. Scrum aparece como un modelo ágil más, no como el ciclo."
            ),
            (
                "El artículo no trata inteligencia artificial ni el prompt. Su valor para esta "
                "intervención es de ancla: Sigemad conserva R, D, I, T y M, y sitúa la IA dentro "
                "de cada fase, en lugar de sustituir el SDLC por un chat continuo o por un marco "
                "ágil pensado solo para equipos humanos."
            ),
        ],
    ),
    (
        "2.3.2 Paladino y Pons (2025): la IA entra al ciclo, con barreras",
        [
            (
                "Paladino y Pons [23] revisan literatura de 2022 a 2025 sobre la integración de "
                "inteligencia artificial en el SDLC. Muestran que la IA ya se usa en planificación, "
                "requisitos, diseño, código, pruebas y mantenimiento, pero que la adopción no es "
                "homogénea: las fases más críticas son planificación, pruebas y despliegue. Las "
                "barreras son técnicas (calidad de datos, integración con sistemas heredados, "
                "código inseguro), organizacionales (capacitación, resistencia al cambio) y éticas "
                "(privacidad, transparencia y explicabilidad)."
            ),
            (
                "El estudio aporta la premisa de que no basta con adoptar un asistente: hace falta "
                "estrategia. No convierte, sin embargo, el prompt en un registro versionado ni "
                "define un ciclo de evaluación con decisión escrita. Esa brecha se opera aquí con "
                "la política mínima de uso de IA (D4), el nivel N2 y la revisión humana antes de "
                "integrar."
            ),
        ],
    ),
    (
        "2.3.3 Hymel: el sprint deja de calibrar el trabajo",
        [
            (
                "Hymel [24] propone un SDLC nativo de IA y el modelo V-Bounce, adaptación del "
                "modelo en V. El argumento central es que, cuando la generación de código se "
                "comprime, el sprint de dos semanas —caja de tiempo pensada para capacidad "
                "humana— deja de dimensionar el trabajo. El esfuerzo se desplaza a requisitos, "
                "arquitectura, diseño y validación continua. El humano pasa de implementador "
                "principal a validador. Cada fase se describe como entrada, generación por IA, "
                "revisión humana, refinamiento, aprobación y captura de conocimiento."
            ),
            (
                "Se cita como white paper de Crowdbotics, no como artículo arbitrado. Su aporte "
                "es teórico y coincide con la unidad de planificación de este caso: el incremento "
                "acotado, no el sprint. Lo que no ofrece es un instrumento auditable "
                "prompt–commit–archivo; esa traza se resuelve con D2, el registro de versiones y "
                "la matriz de doble entrada."
            ),
        ],
    ),
    (
        "2.3.4 Wang et al.: la calidad del prompt condiciona la seguridad del código",
        [
            (
                "Wang et al. [7] muestran, con un banco de tareas en Python alineado a CWE, que "
                "un prompt benévolo pero pobre —poco claro, incompleto o lógicamente inconsistente— "
                "aumenta de forma marcada la probabilidad de código inseguro. Técnicas como "
                "cadena de pensamiento y autocorrección reducen ese riesgo. El hallazgo desplaza "
                "el problema de “el modelo es inseguro” hacia “el pedido al modelo es un control "
                "de calidad”."
            ),
            (
                "El experimento no está hecho en PHP ni en un dominio municipal; no se extrapola "
                "el porcentaje de defectos a Sigemad. Sí se retiene la relación calidad del prompt "
                "↔ seguridad, que justifica la anatomía D1 (objetivo, restricciones y bloque «No "
                "hacer»), la prohibición de secretos en el prompt y la revisión humana del "
                "Definition of Done. El paper no versiona prompts ni los ata a un commit."
            ),
        ],
    ),
    (
        "2.3.5 Esposito et al. (2026): GenAI en arquitectura y salidas sin evaluación rigurosa",
        [
            (
                "Esposito et al. [25] publican en Journal of Systems and Software una revisión "
                "multivocal sobre GenAI en arquitectura de software. Encuentran uso concentrado "
                "en las etapas iniciales del ciclo de arquitectura (de requisitos a arquitectura y "
                "de arquitectura a código), con GPT, few-shot y RAG, sobre todo en monolitos y "
                "microservicios. El hallazgo crítico para esta intervención es que las salidas "
                "casi nunca se someten a prueba rigurosa, y que persisten alucinaciones, huecos "
                "éticos y ausencia de marcos de evaluación."
            ),
            (
                "El artículo no cubre el ciclo completo ni un repositorio de prompts. Sí justifica "
                "que un ADR o un esquema generado por el agente no se aceptan por el hecho de "
                "haber sido generados: en Sigemad esa evaluación es D2 (métrica, diagnóstico, "
                "refinamiento y decisión Aprobado / Refinado / Rechazado / Dividir / Superado) y "
                "se aplica de forma explícita a los entregables de diseño (D-001 a D-006, ADR-001 "
                "y ADR-002)."
            ),
        ],
    ),
]

SINTESIS_H = "2.3.6 Síntesis"
SINTESIS = [
    (
        "Ninguna de las cinco fuentes, por sí sola, describe el Prompt-Centered SDLC. Yas [22] "
        "da las fases y niega un modelo único; Paladino y Pons [23] documentan que la IA ya "
        "atraviesa el ciclo y que faltan gobernanza y ética; Hymel [24] explica por qué el sprint "
        "deja de ser la unidad útil; Wang et al. [7] prueban que la norma del prompt es un "
        "control de seguridad; Esposito et al. [25] muestran que el diseño asistido por GenAI "
        "se publica con poca evaluación. La Tabla 4 concentra esas brechas."
    ),
    (
        "Lo que esta intervención añade, y ninguna de esas fuentes opera como sistema, es el "
        "cierre conjunto: prompt versionado (D1), decisión D2, mapa por fase R/D/I/T/M (D3), "
        "política de IA en N2 (D4) y traza prompt–commit–archivo (D5). Otras lecturas del corpus "
        "(Shrivastava et al., 2025; Kirchner et al., 2025; Bandara et al. [2]) se usan más "
        "adelante: ingeniería de prompts y repositorio, gestión del conocimiento, y Agentsway "
        "como metodología agentic afín que no se aplicó en este caso."
    ),
]

TABLE_ROWS = [
    [
        "Yas, Alazzawi y Rahmatullah (2023) [22]",
        "Fases estables del SDLC; Scrum es un modelo ágil más, no el ciclo",
        "IA y prompt como artefacto",
        "Se conservan R/D/I/T/M; la IA entra dentro de cada fase",
    ],
    [
        "Paladino y Pons (2025) [23]",
        "IA transversal al SDLC; barreras técnicas, organizacionales y éticas",
        "Prompt versionado, ciclo D2 y traza Git",
        "Política de IA (D4), nivel N2 y revisión humana antes de integrar",
    ],
    [
        "Hymel (V-Bounce) [24]",
        "El sprint deja de calibrar; humano como validador; ciclo entrada–generación–revisión",
        "Instrumento auditable prompt–commit–archivo (es white paper, no revista)",
        "Unidad de planificación = incremento acotado + registro D2",
    ],
    [
        "Wang et al. [7]",
        "Prompt poco normativo aumenta código inseguro; CoT y autocorrección mitigan",
        "Versionado del prompt y caso PHP municipal",
        "Anatomía D1, bloque «No hacer» y secretos fuera del prompt",
    ],
    [
        "Esposito et al. (2026) [25]",
        "GenAI en arquitectura; few-shot/RAG; salidas casi nunca se evalúan con rigor",
        "Ciclo completo, repositorio de prompts y matriz de funciones",
        "D2 sobre entregables de diseño (ADR, esquema) antes de implementar",
    ],
]

NEW_REFS = [
    (
        "[22]\tQ. M. Yas, A. Alazzawi y B. Rahmatullah, “A Comprehensive Review of Software "
        "Development Life Cycle Methodologies: Pros, Cons, and Future Directions,” Iraqi Journal "
        "for Computer Science and Mathematics, vol. 4, núm. 4, art. 14, 2023, doi: "
        "10.52866/ijcsm.2023.04.04.014."
    ),
    (
        "[23]\tM. Paladino y C. Pons, “From Obstacles to Strategies: Integrating Artificial "
        "Intelligence into the Software Development Life Cycle,” en ASSE, Argentine Symposium "
        "on Software Engineering, 2025."
    ),
    (
        "[24]\tC. Hymel, “The AI-Native Software Development Lifecycle: A Theoretical and "
        "Practical New Methodology,” Crowdbotics, white paper, 2024."
    ),
    (
        "[25]\tM. Esposito, X. Li, S. Moreschini, N. Ahmad, T. Cerny, K. Vaidhyanathan, "
        "V. Lenarduzzi y D. Taibi, “Generative AI for software architecture. Applications, "
        "challenges, and future directions,” The Journal of Systems & Software, vol. 231, "
        "art. 112607, 2026, doi: 10.1016/j.jss.2025.112607."
    ),
]


def fill_table(doc):
    table = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Fuente":
            table = t
            break
    if table is None:
        raise RuntimeError("No se encontró la Tabla 4 (encabezado Fuente).")
    while len(table.rows) < 6:
        table.add_row()
    for i, row_data in enumerate(TABLE_ROWS, start=1):
        for j, text in enumerate(row_data):
            set_cell_text(table.rows[i].cells[j], text, size=8)


def insert_23_body(doc):
    heading = find_body_heading(doc, "2.3 Posicionamiento frente a la literatura de IA en el SDLC")
    if heading is None:
        raise RuntimeError("No se encontró el apartado 2.3.")
    if has_text(doc, "2.3.1 Yas, Alazzawi y Rahmatullah"):
        return
    el = heading._element
    p = insert_para(el, INTRO, doc)
    el = p._element
    for title, paras in SECTIONS:
        h = insert_para(el, title, doc, bold=True, first_indent=0, after=6, justify=False)
        el = h._element
        for text in paras:
            b = insert_para(el, text, doc)
            el = b._element


def insert_sintesis(doc):
    if has_text(doc, "2.3.6 Síntesis"):
        return
    matches = [p for p in doc.paragraphs if p.text.strip().startswith("Tabla 4. Literatura")]
    if not matches:
        raise RuntimeError("No se encontró el título de la Tabla 4.")
    cap = matches[-1]
    set_runs_text(
        cap,
        "Tabla 4. Literatura de IA en el SDLC: cinco fuentes (aporta, no cubre y brecha).",
        size=11,
        italic=True,
        justify=False,
        first_indent=0,
        after=8,
        center=True,
    )
    el = cap._element
    h = insert_para(el, SINTESIS_H, doc, bold=True, first_indent=0, after=6, justify=False)
    el = h._element
    for text in SINTESIS:
        b = insert_para(el, text, doc)
        el = b._element


def insert_refs(doc):
    if has_text(doc, "[22]\tQ. M. Yas"):
        return
    last = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("[21]"):
            last = p
    if last is None:
        raise RuntimeError("No se encontró [21].")
    el = last._element
    for text in NEW_REFS:
        p = insert_para(el, text, doc, first_indent=0, after=4, justify=False)
        el = p._element


def main():
    doc = Document(str(DOC))
    insert_23_body(doc)
    fill_table(doc)
    insert_sintesis(doc)
    insert_refs(doc)
    doc.save(str(DOC))
    print("OK 2.3 cinco fuentes:", DOC.name)


if __name__ == "__main__":
    main()

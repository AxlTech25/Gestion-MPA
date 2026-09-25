# -*- coding: utf-8 -*-
"""Amplía D2, matriz de doble entrada, seguridad/monitoreo y cumplimiento de política. No toca el apartado 8."""
import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent
DOC = BASE / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
IMG = BASE / "_figuras_anexo7"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)


def _font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_body(p, size=12, first_indent=0.75, after=8, justify=True, center=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    )
    pf = p.paragraph_format
    pf.space_before = Pt(0 if not center else 4)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0 if center else 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0 if center else first_indent)


def set_runs_text(p, text, size=12, bold=False, italic=False, justify=True, first_indent=0.75, after=8, center=False):
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size, bold=bold, italic=italic)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size, bold=bold, italic=italic)
    style_body(p, size=size, first_indent=first_indent, after=after, justify=justify, center=center)


def insert_after(element, doc):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    return Paragraph(new_p, doc)


def insert_body_after(element, text, doc, size=12, first_indent=0.75, after=8):
    p = insert_after(element, doc)
    run = p.add_run(text)
    _font(run, size=size)
    style_body(p, size=size, first_indent=first_indent, after=after)
    return p


def find_body_heading(doc, text):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    return matches[-1] if matches else None


def has_text(doc, needle):
    return any(needle in p.text for p in doc.paragraphs)


def first_body_after(heading):
    return None


def is_next_heading(t):
    if t.startswith(("Tabla ", "Figura ", "Gráfico")):
        return True
    if re.match(r"^\d+(\.\d+)*\s", t):
        return True
    return False


def fill_section(doc, heading_text, paragraphs, replace_startswith=None):
    heading = find_body_heading(doc, heading_text)
    if heading is None:
        raise RuntimeError("No se encontró " + heading_text)
    body = None
    seen = False
    for p in doc.paragraphs:
        if p is heading:
            seen = True
            continue
        if not seen:
            continue
        t = p.text.strip()
        if not t:
            continue
        if is_next_heading(t):
            break
        body = p
        break
    if body is None:
        if has_text(doc, paragraphs[0][:50]):
            return heading
        el = heading._element
        for text in paragraphs:
            p = insert_body_after(el, text, doc)
            el = p._element
        return heading
    if replace_startswith is None or body.text.strip().startswith(replace_startswith) or body.text.strip().startswith(paragraphs[0][:40]):
        set_runs_text(body, paragraphs[0])
    if len(paragraphs) == 1:
        return heading
    if has_text(doc, paragraphs[1][:52]):
        return heading
    el = body._element
    for text in paragraphs[1:]:
        p = insert_body_after(el, text, doc)
        el = p._element
    return heading


def insert_after_caption(doc, caption_start, paragraphs, unique_needle):
    if has_text(doc, unique_needle):
        return
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(caption_start)]
    if not matches:
        raise RuntimeError("No se encontró título " + caption_start)
    el = matches[-1]._element
    for text in paragraphs:
        p = insert_body_after(el, text, doc)
        el = p._element


def replace_picture_before_caption(doc, caption_start, image_path, width_cm=15.8):
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        if not p.text.strip().startswith(caption_start):
            continue
        # use last caption (body, not TOC)
    last_i = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith(caption_start):
            last_i = i
    if last_i is None:
        return False
    for j in range(last_i - 1, max(-1, last_i - 6), -1):
        q = paras[j]
        drawings = q._element.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
        )
        if drawings or not q.text.strip():
            for child in list(q._element):
                tag = child.tag
                if tag.endswith("}r") or tag.endswith("}hyperlink"):
                    q._element.remove(child)
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.first_line_indent = Cm(0)
            q.paragraph_format.space_before = Pt(6)
            q.paragraph_format.space_after = Pt(2)
            run = q.add_run()
            run.add_picture(str(image_path), width=Cm(width_cm))
            return True
    return False


def _new_fig(w, h):
    plt.rcParams["font.family"] = "Times New Roman"
    fig, ax = plt.subplots(figsize=(w, h), dpi=170)
    fig.patch.set_facecolor("white")
    ax.axis("off")
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)
    return fig, ax


def _round(ax, x, y, w, h, facecolor, edge="#1F3A5F", lw=1.25):
    ax.add_patch(
        FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            linewidth=lw, edgecolor=edge, facecolor=facecolor,
        )
    )


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1), (x2, y2),
            arrowstyle="-|>", mutation_scale=11, linewidth=1.2, color="#1F3A5F",
        )
    )


def draw_d2():
    IMG.mkdir(exist_ok=True)
    fig, ax = _new_fig(11.3, 6.15)
    ax.text(5.65, 5.92, "Cómo se opera el ciclo D2 en este caso",
            ha="center", va="center", fontsize=13.5, fontweight="bold", color="#1F3A5F")
    ax.text(5.65, 5.60, "El humano mide la salida del agente, diagnostica el desvío, refina el prompt y deja una decisión escrita.",
            ha="center", va="center", fontsize=8.3, color="#4A5568")

    steps = [
        (0.25, "1. Medición", "¿El artefacto cumple\nlo pedido?", "Tests, HU, plantilla D1\ny revisión humana", "#C5DDD4"),
        (3.05, "2. Diagnóstico", "¿Qué falló en el\nprompt o en el código?", "Tarea ancha, restricción\nausente, alucinación", "#D6E4F0"),
        (5.85, "3. Refinamiento", "Se ajusta el prompt\no se parte la tarea", "v1.1 / v2, o Dividir\nen varios I-*", "#B8CDE0"),
        (8.65, "4. Decisión", "Queda escrita en el\nregistro del prompt", "Aprobado, Refinado,\nRechazado, Superado", "#F3E4C8"),
    ]
    bw, bh, y = 2.40, 2.05, 3.20
    for x, title, q, how, col in steps:
        _round(ax, x, y, bw, bh, col)
        ax.text(x + bw / 2, y + 1.72, title, ha="center", va="center",
                fontsize=9.4, fontweight="bold", color="#1F3A5F")
        ax.text(x + bw / 2, y + 1.18, q, ha="center", va="center",
                fontsize=7.8, color="#1A1A1A", linespacing=1.2)
        ax.text(x + bw / 2, y + 0.42, how, ha="center", va="center",
                fontsize=7.2, color="#4A5568", linespacing=1.2)
    for x in (2.65, 5.45, 8.25):
        _arrow(ax, x, y + bh / 2, x + 0.40, y + bh / 2)

    ax.add_patch(Rectangle((0.25, 0.18), 10.80, 2.78, facecolor="#F4F7FB", edgecolor="#1F3A5F", lw=1.05))
    ax.text(0.45, 2.68, "Qué significa cada decisión en Sigemad", ha="left", va="center",
            fontsize=8.6, fontweight="bold", color="#1F3A5F")
    outcomes = [
        ("Aprobado", "Se cierra el registro. Ejemplo: I-002 inventario."),
        ("Refinado / Dividir", "Se vuelve a medir. Ejemplo: macro I-001 partido en I-001…I-009."),
        ("Rechazado", "No se integra. Se abre un prompt nuevo."),
        ("Superado", "El archivo se conserva, pero ya no se cita. Ejemplo: I-002-ML."),
    ]
    yy = 2.28
    for title, desc in outcomes:
        ax.text(0.55, yy, "•  " + title + "  —  " + desc, ha="left", va="center",
                fontsize=7.7, color="#1A1A1A")
        yy -= 0.48
    path = IMG / "fig_a7_9_d2.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white", pad_inches=0.08)
    plt.close()
    return path


def draw_matriz():
    fig, ax = _new_fig(11.3, 6.45)
    ax.text(5.65, 6.20, "Cómo se lee la matriz de doble entrada",
            ha="center", va="center", fontsize=13.5, fontweight="bold", color="#1F3A5F")
    ax.text(5.65, 5.88, "Cada fila responde cuatro preguntas. El ejemplo es el inventario de equipos (F-003).",
            ha="center", va="center", fontsize=8.3, color="#4A5568")

    boxes = [
        (0.22, "1. Función", "¿Qué parte del\nsistema miro?", "Inventario\n(F-003)", "#C5DDD4"),
        (3.02, "2. Prompt", "¿Qué se le pidió\na la IA?", "I-002\nImplementación", "#D6E4F0"),
        (5.82, "3. Commit", "¿En qué cambio\nquedó el código?", "GitHub\nf086a63", "#B8CDE0"),
        (8.62, "4. Archivo", "¿Qué archivo abro\npara auditar?", "EquipoController.php\nen ese SHA", "#F3E4C8"),
    ]
    bw, bh, y = 2.46, 2.22, 3.28
    for x, title, q, ans, col in boxes:
        _round(ax, x, y, bw, bh, col)
        ax.text(x + bw / 2, y + 1.90, title, ha="center", va="center",
                fontsize=9.6, fontweight="bold", color="#1F3A5F")
        ax.text(x + bw / 2, y + 1.28, q, ha="center", va="center",
                fontsize=7.6, color="#4A5568", linespacing=1.2)
        ax.text(x + bw / 2, y + 0.48, ans, ha="center", va="center",
                fontsize=8.4, fontweight="bold", color="#1A1A1A", linespacing=1.15)
    for x in (2.68, 5.48, 8.28):
        _arrow(ax, x, y + bh / 2, x + 0.34, y + bh / 2)

    ax.add_patch(Rectangle((0.22, 0.18), 10.86, 2.88, facecolor="#F4F7FB", edgecolor="#1F3A5F", lw=1.05))
    ax.text(0.42, 2.76, "Por qué se llama de doble entrada", ha="left", va="center",
            fontsize=8.8, fontweight="bold", color="#1F3A5F")
    lines = [
        "Una entrada es la función del producto (F-001 autenticación, F-003 inventario, F-007 ML…).",
        "La otra entrada es la evidencia de cómo se construyó: prompt + commit + archivo.",
        "Si las tres piezas se pueden abrir en GitHub, un tercero puede auditar sin entrar al chat.",
        "Ejemplo: F-003 → I-002 → commit f086a63 → EquipoController.php. El Excel completo es el Apéndice A.",
    ]
    yy = 2.32
    for line in lines:
        ax.text(0.50, yy, "•  " + line, ha="left", va="center", fontsize=7.7, color="#1A1A1A")
        yy -= 0.48
    path = IMG / "fig_a7_11_cadena.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white", pad_inches=0.08)
    plt.close()
    return path


D2_61 = [
    (
        "D2 es el ciclo con el que se evalúa cada salida del agente antes de darla por cerrada. "
        "No basta con que el modelo haya generado texto o código: un humano mide, diagnostica, "
        "refina si hace falta y deja una decisión escrita. Ese ciclo vive en el bloque «Evaluación D2» "
        "de cada registro de prompt y se consolida en prompts/gobernanza/registro_metricas.md. "
        "La Figura 12 muestra los cuatro pasos y el significado de cada decisión en este caso."
    ),
    (
        "La medición compara el artefacto con criterios observables: historias y criterios de aceptación, "
        "completitud de la plantilla D1, pruebas relevantes (plan T-001 y runners T-002 a T-004) y ausencia "
        "de secretos. El diagnóstico, si hay desvío, pregunta qué falló en el prompt (tarea demasiado ancha, "
        "restricción omitida, formato incorrecto) o en el código (por ejemplo, autorización solo en la "
        "interfaz). El refinamiento produce una versión v1.1 o v2 del mismo registro, o la decisión de "
        "Dividir la tarea en varios prompts. La decisión final queda en el registro: Aprobado, Refinado, "
        "Rechazado, Dividir o Superado."
    ),
    (
        "En Sigemad el ciclo se aplicó de forma mixta, y eso se declara. Varios I-* se reconstruyeron a "
        "posteriori y sus iteraciones de origen no se midieron en caliente (Tabla 15). Sí hubo ciclos "
        "documentados en caliente o con evidencia de más de una pasada: I-007 requirió al menos dos "
        "iteraciones (el lote del modelo no aceptaba una lista vacía y el proxy HTTP); I-009 nació de un "
        "diagnóstico de seguridad sobre I-005; R-004 y D-006 se ejecutaron en la oleada 3 con medición "
        "explícita. El criterio de parada es no seguir iterando sin dejar decisión D2."
    ),
]

D2_FIG = [
    (
        "La Figura 12 se lee de izquierda a derecha. El responsable no le pide al modelo que «se autoevalúe»: "
        "mide con pruebas y revisión, diagnostica el desvío, cambia el prompt o lo parte, y escribe la "
        "decisión. Aprobado cierra el registro. Refinado o Dividir vuelve a la medición. Rechazado impide "
        "integrar. Superado conserva el archivo histórico, pero deja de citarse en commits nuevos."
    ),
]

D2_62 = [
    (
        "El registro de versiones no es el historial del chat. Es la ficha de cada prompt: código, fase, "
        "versión (v1, v1.1, v2), número de iteraciones cuando se conoció, medición, decisión y lección "
        "reutilizable. Cada archivo en prompts/ tiene ese bloque; el tablero vivo es registro_metricas.md. "
        "La Tabla 15 extrae cinco filas representativas para el cuerpo del anexo; el resto queda en el "
        "Apéndice B y en el registro de métricas."
    ),
    (
        "Cómo se lee la Tabla 15. Código identifica el prompt (R-004, I-007). Fase sitúa el registro en el "
        "ciclo de vida. Iteraciones indica cuántas pasadas hubo; «No medido» significa que no se contó en "
        "caliente, no que el trabajo fuera trivial. Medición D2 es la evidencia con la que se juzgó el "
        "output (hallazgos, ADR, historias, métrica de modelo). Decisión es el cierre. Lección es lo que "
        "el siguiente prompt debe heredar. Los macros Superados no se borran: dejan de ser vigentes."
    ),
]

D2_63 = [
    (
        "El caso de refinamiento que mejor muestra D2 es el macro-prompt I-001 (API, autenticación e "
        "inventario en un solo registro) y el alias I-002-ML. La medición falló el criterio de granularidad: "
        "un humano no podía aprobar ni citar ese identificador en un commit. El diagnóstico fue «tarea "
        "demasiado ancha». La decisión D2 fue Superado, no borrar. El refinamiento consistió en Dividir: "
        "I-001 a I-009, un prompt por incremento. Los vigentes quedaron Aprobados; los macros permanecen "
        "en el repositorio para no reescribir la historia. El antes y el después se resumen en la Tabla 10b."
    ),
    (
        "Un segundo patrón, más fino, es I-009. I-005 entregó configuración, pero la autorización quedó "
        "solo en la interfaz. La medición D2 (criterio de seguridad y DoD) no cerró el hueco. El diagnóstico "
        "exigió un prompt nuevo, no un parche silencioso en el mismo archivo. I-009 incorporó requireRole "
        "y se aprobó como parche 0.9.1. En D2, un hallazgo de seguridad no se «arregla en el chat»: se "
        "versiona como registro propio."
    ),
]

MATRIZ_71 = [
    (
        "La matriz de doble entrada no es una tabla decorativa: es el instrumento con el que un tercero "
        "puede reconstruir, sin abrir el chat, cómo se construyó cada función del software. Una entrada "
        "es la función del producto (F-001 autenticación, F-003 inventario, F-007 aprendizaje automático, "
        "y así sucesivamente). La otra entrada es la evidencia de construcción: el prompt versionado, el "
        "commit de GitHub y el archivo tal como quedó en ese identificador SHA. Si esas tres piezas se "
        "pueden abrir, hay trazabilidad; si falta una, el incremento no es auditable."
    ),
    (
        "La Figura 13 explica la lectura con un ejemplo real del caso. Para el inventario (F-003) la "
        "pregunta «qué se le pidió a la IA» se responde con I-002; «dónde quedó el código» con el commit "
        "f086a63; «qué archivo abro» con EquipoController.php en ese SHA. El mismo esquema se aplica a "
        "las demás funciones. Los SHA no se repiten en este apartado: constan en la Tabla 14 y el Excel "
        "completo es el Apéndice A."
    ),
]

MATRIZ_FIG = [
    (
        "La Figura 13 se lee como cuatro preguntas en cadena, no como un diagrama técnico. Función, "
        "prompt, commit y archivo deben coincidir. Un mismo commit puede anclar dos funciones (por "
        "ejemplo F-003 y F-008 en f086a63), pero cada una conserva su prompt. I-008 e I-009 son "
        "evoluciones posteriores: no borran el SHA de origen."
    ),
]

MATRIZ_73 = [
    (
        "En el Excel, cada fila agrupa: identificación de la función (código F-00N y módulo); impacto "
        "sobre el ciclo de vida; cinco capacidades DevOps (control de versiones, integración continua, "
        "despliegue continuo, pruebas y monitoreo) con estado Sí, Parcial o No; datos del prompt "
        "(técnica, iteraciones, decisión D2); ruta del archivo y SHA; prueba asociada; y observaciones. "
        "En el cuerpo del anexo no se reproduce esa grilla completa. La Tabla 16 muestra solo el extracto "
        "DevOps; la Tabla 14, los SHA."
    ),
]

MATRIZ_74 = [
    (
        "Se aplicó a F-001 a F-008 y a las evoluciones I-008 e I-009. El control de versiones está "
        "implantado en GitHub. La integración continua y el despliegue continuo no están implantados: "
        "el caso aún no opera un ambiente de producción institucional con pipeline. Las pruebas unitarias "
        "existen por entorno (PHPUnit, Vitest, pytest), con cobertura de rama aún no medida. El monitoreo "
        "de aplicaciones en producción no aplica en este momento; el diseño de degradación del componente "
        "predictivo (ADR-002) queda como control preparado. El extracto está en la Tabla 16."
    ),
]

MATRIZ_75 = [
    (
        "Para un cambio futuro la matriz se usa así. Primero se ubica la función (por ejemplo inventario). "
        "Después se abre el prompt de origen y se leen medición, decisión y lección. Luego se abre el "
        "archivo en el commit de ancla, no necesariamente en la punta de la rama. Se ejecuta la prueba "
        "asociada. Si el cambio usa IA, se versiona un prompt nuevo o una vN+1, se cita en el commit y se "
        "actualizan iteraciones y estado DevOps en el Excel. Sin ese paso, la doble entrada se rompe."
    ),
]

SEG_91 = [
    (
        "La seguridad se gestiona en dos momentos, ambos con evidencia. Antes de ejecutar el prompt, en la "
        "anatomía D1: el bloque No hacer prohíbe secretos, datos personales reales, reescritura total de V1 "
        "y dependencias no autorizadas; las restricciones anclan el modelo al SQL y a los ADR existentes; "
        "la tarea se acota a un incremento para que el agente no invente módulos. Después de generar, el "
        "control posterior es revisión humana, pruebas, diff en GitHub y, cuando el hallazgo es de "
        "autorización, un prompt nuevo (I-009). El cruce riesgo–guardrail–control está en la Tabla 18."
    ),
    (
        "La gobernanza no es un comité aparte: es la política v1.3 (apartado 4.1 y Tabla 6), el catálogo "
        "de prompts, el registro D2, los ADR y el DoD. Cursor Agent genera; el tesista define el prompt, "
        "revisa, prueba e integra. No hay encadenamiento de agentes sin revisión por paso. Los secretos "
        "(local.php, JWT, contraseñas) permanecen fuera de Git, con respaldo en M-004. Las decisiones de "
        "arquitectura no se dejan en un comentario del chat: ADR-001 (stack) y ADR-002 (degradación del "
        "ML si el servicio no responde)."
    ),
]

SEG_92 = [
    (
        "El sistema aún no está en producción institucional. El monitoreo que se describe aquí es el de "
        "desarrollo y de preparación para el despliegue, no un monitor de aplicaciones en servidor público. "
        "En cada incremento se revisan: resultados de PHPUnit, Vitest y pytest; el plan funcional T-001; "
        "errores de autenticación en el flujo local; y el comportamiento de degradación del componente "
        "predictivo según ADR-002 (la interfaz no debe quedar en blanco si FastAPI no responde). No se "
        "implantó APM ni bitácora centralizada de producción."
    ),
    (
        "Cuando exista publicación, el control previsto es M-001: pruebas de humo, comprobación de que "
        "los secretos no viajaron al repositorio y plan de retorno (M-004). Hasta entonces, la frecuencia "
        "de revisión es la del propio incremento y la de cada decisión D2, no un tablero 24/7. Esa "
        "limitación se declara para no atribuir al caso una operación que todavía no tiene."
    ),
]

SEG_93 = [
    (
        "El cumplimiento de la política v1.3 se actualiza con la evidencia del propio anexo, no con una "
        "declaración genérica. Las reglas 1 a 5 y 8 de la Tabla 6 están cubiertas: hay revisor en cada "
        "registro, no hay secretos en el repositorio, los prompts viven en GitHub, existen ADR-001 y "
        "ADR-002, se ejecutan pruebas por runner y los macros Superados ya no se citan. La regla 6 "
        "(poder mantener el código sin el modelo) se cumple en la medida en que un solo desarrollador-revisor "
        "puede explicar los incrementos; no hay un segundo par independiente."
    ),
    (
        "La regla 7 se cumple de forma honesta y parcial: los registros distinguen Reconstruido y Ejecutado, "
        "pero una parte de los I-* se documentó a posteriori. No hay todavía una auditoría periódica formal "
        "de la política (apartado 12). En síntesis: la política está operativa como disciplina de trabajo "
        "y como evidencia versionada; no está certificada ni auditada por un tercero, y el caso no debe "
        "presentarse como un despliegue en producción ya gobernado de punta a punta."
    ),
]


def main():
    img_d2 = draw_d2()
    img_mx = draw_matriz()
    print("PNG", img_d2)
    print("PNG", img_mx)

    doc = Document(str(DOC))

    fill_section(doc, "6.1 Ciclo D2 aplicado", D2_61, "El ciclo D2 consiste")
    insert_after_caption(doc, "Figura 12.", D2_FIG, "La Figura 12 se lee de izquierda a derecha")

    fill_section(doc, "6.2 Registro de versiones", D2_62)
    fill_section(doc, "6.3 Ejemplo de refinamiento", D2_63, "El refinamiento documentado")

    fill_section(doc, "7.1 Propósito de la matriz", MATRIZ_71, "La matriz de doble entrada V3 conecta")
    insert_after_caption(doc, "Figura 13.", MATRIZ_FIG, "La Figura 13 se lee como cuatro preguntas")
    fill_section(doc, "7.3 Estructura de la matriz", MATRIZ_73, "La matriz agrupa identificación")
    fill_section(doc, "7.4 Aplicación de la matriz en este caso", MATRIZ_74, "Se aplicó a las funciones F-001")
    fill_section(doc, "7.5 Uso de la matriz para mantenimiento", MATRIZ_75, "Para un cambio futuro se propone")

    fill_section(doc, "9.1 Guardrails en los prompts", SEG_91, "Las restricciones de seguridad se escribieron")
    fill_section(doc, "9.2 Monitoreo", SEG_92, "En el ambiente de producción se vigilan")
    fill_section(doc, "9.3 Cumplimiento de la política de uso de IA", SEG_93, "La política de uso de IA se cumplió")

    cap12 = [p for p in doc.paragraphs if p.text.strip().startswith("Figura 12.")]
    if cap12:
        set_runs_text(
            cap12[-1],
            "Figura 12. Ciclo D2 aplicado: medición, diagnóstico, refinamiento y decisión, con el significado de cada cierre en Sigemad.",
            size=11, italic=True, center=True, first_indent=0, after=8,
        )
    cap13 = [p for p in doc.paragraphs if p.text.strip().startswith("Figura 13.")]
    if cap13:
        set_runs_text(
            cap13[-1],
            "Figura 13. Lectura de la matriz de doble entrada: función, prompt, commit y archivo (ejemplo F-003 inventario).",
            size=11, italic=True, center=True, first_indent=0, after=8,
        )

    ok12 = replace_picture_before_caption(doc, "Figura 12.", img_d2, 16.0)
    ok13 = replace_picture_before_caption(doc, "Figura 13.", img_mx, 16.0)
    print("Figuras", ok12, ok13)

    doc.save(str(DOC))
    print("OK", DOC)


if __name__ == "__main__":
    main()

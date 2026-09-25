# -*- coding: utf-8 -*-
"""Genera la plantilla Word del Anexo 7 (guía + tablas y figuras de ejemplo)."""
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
IMG = BASE / "_figuras_anexo7"
OUT = BASE / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x1A, 0x5F, 0x7A)
GRAY = RGBColor(0x4A, 0x55, 0x68)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Times New Roman"


def _set_run_font(run, size=12, bold=False, italic=False, color=BLACK, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, val in kwargs.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), val.get("val", "single"))
        el.set(qn("w:sz"), val.get("sz", "4"))
        el.set(qn("w:color"), val.get("color", "1F3A5F"))


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_paragraph_spacing(p, before=0, after=8, line=1.5, first_line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    _set_run_font(run, size=10, color=GRAY)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Anexo 07  ·  Plantilla guía  ·  Prompt-Centered SDLC v1.2")
    _set_run_font(run, size=9, italic=True, color=GRAY)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = fp.add_run("Sigemad MPA  ·  Municipalidad Provincial de Acobamba")
    _set_run_font(left, size=9, color=GRAY)
    mid = fp.add_run("     —     ")
    _set_run_font(mid, size=9, color=GRAY)
    add_page_number(fp)


def add_heading_custom(doc, text, level):
    sizes = {0: 20, 1: 15, 2: 13, 3: 12}
    colors = {0: NAVY, 1: NAVY, 2: TEAL, 3: NAVY}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    before = {0: 6, 1: 16, 2: 12, 3: 8}[level]
    after = {0: 10, 1: 8, 2: 6, 3: 4}[level]
    set_paragraph_spacing(p, before=before, after=after, line=1.15, first_line=0)
    run = p.add_run(text)
    _set_run_font(run, size=sizes[level], bold=True, color=colors[level])
    if level in (0, 1):
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12" if level == 0 else "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1F3A5F" if level == 0 else "1A5F7A")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_body(doc, text, first_indent=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=8, line=1.5, first_line=first_indent)
    run = p.add_run(text)
    _set_run_font(run, size=12)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=12, line=1.0, first_line=0)
    run = p.add_run(text)
    _set_run_font(run, size=10, italic=True, color=GRAY)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=2, after=10, line=1.15, first_line=0)
    label = p.add_run("Cómo usarla. ")
    _set_run_font(label, size=10, italic=True, bold=True, color=GRAY)
    run = p.add_run(text)
    _set_run_font(run, size=10, italic=True, color=GRAY)
    return p


def add_callout(doc, text, fill="F2F2F2", lead="Redactar. "):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    shade_cell(cell, fill)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "8", "color": "8AA0B8"},
        bottom={"val": "single", "sz": "8", "color": "8AA0B8"},
        left={"val": "single", "sz": "8", "color": "8AA0B8"},
        right={"val": "single", "sz": "8", "color": "8AA0B8"},
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    lab = p.add_run(lead)
    _set_run_font(lab, size=10, bold=True, italic=True, color=TEAL if fill != "F2F2F2" else GRAY)
    run = p.add_run(text)
    _set_run_font(run, size=10, italic=True, color=GRAY)
    doc.add_paragraph()
    return table


def add_table(doc, headers, rows, col_widths=None, font_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = Cm(15.5)
    if col_widths is None:
        col_widths = [usable / len(headers)] * len(headers)
    for i, w in enumerate(col_widths):
        for cell in table.columns[i].cells:
            cell.width = w
    border = {"val": "single", "sz": "4", "color": "8AA0B8"}
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "1F3A5F")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        run = cell.paragraphs[0].add_run(h)
        _set_run_font(run, size=font_size, bold=True, color=WHITE)
        set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    for r_idx, row in enumerate(rows):
        prevent_row_split(table.rows[r_idx + 1])
        fill = "F4F7FB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            shade_cell(cell, fill)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            run = cell.paragraphs[0].add_run(str(value))
            _set_run_font(run, size=font_size)
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    return table


def add_picture_centered(doc, path, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=6, after=2, line=1.0, first_line=0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def _box(ax, x, y, w, h, text, facecolor, edge="#1F3A5F", fontsize=8.5, sub=None):
    ax.add_patch(
        FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            linewidth=1.3, edgecolor=edge, facecolor=facecolor,
        )
    )
    ax.text(
        x + w / 2, y + h / 2 + (0.1 if sub else 0), text,
        ha="center", va="center", fontsize=fontsize, fontweight="bold",
        color="#1A1A1A", wrap=True,
    )
    if sub:
        ax.text(
            x + w / 2, y + h / 2 - 0.2, sub,
            ha="center", va="center", fontsize=6.8, color="#4A5568",
        )


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1), (x2, y2),
            arrowstyle="-|>", mutation_scale=12, linewidth=1.3, color="#1F3A5F",
        )
    )


def _new_fig(w=10.2, h=3.6):
    fig, ax = plt.subplots(figsize=(w, h), dpi=160)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    return fig, ax


def draw_all_figures():
    IMG.mkdir(exist_ok=True)

    fig, ax = _new_fig(10.2, 2.8)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 2.8)
    labels = [
        (0.2, "Prompt\nversionado"),
        (2.2, "Artefacto"),
        (4.2, "Métrica D2"),
        (6.2, "Decisión"),
        (8.2, "Siguiente\nfase"),
    ]
    for x, t in labels:
        _box(ax, x, 1.0, 1.7, 1.2, t, "#D6E4F0")
    for x in (1.9, 3.9, 5.9, 7.9):
        _arrow(ax, x, 1.6, x + 0.3, 1.6)
    ax.annotate(
        "si se refina",
        xy=(1.05, 1.0), xytext=(6.9, 0.25),
        fontsize=7.5, color="#1A5F7A",
        arrowprops=dict(arrowstyle="-|>", color="#1A5F7A", lw=1.0, connectionstyle="arc3,rad=-0.25"),
    )
    fig.savefig(IMG / "fig_a7_1_cadena.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 4.4)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 4.4)
    dims = [
        (0.15, "D1 Anatomía", "Plantilla y registros\nR / D / I / T / M"),
        (2.15, "D2 Evaluación", "registro_metricas.md"),
        (4.15, "D3 Mapa fase", "catálogo: 29 vigentes"),
        (6.15, "D4 Gobernanza", "politicas_uso_ia.md"),
        (8.15, "D5 Trazabilidad", "Matriz V3 Amaro–Jiang"),
    ]
    for x, t, s in dims:
        _box(ax, x, 2.5, 1.9, 1.5, t, "#B8CDE0")
        _arrow(ax, x + 0.95, 2.5, x + 0.95, 1.85)
        _box(ax, x, 0.25, 1.9, 1.5, s, "#F4F7FB", fontsize=7.2)
    fig.savefig(IMG / "fig_a7_2_d1d5.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 5.2)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 5.2)
    _box(ax, 0.3, 4.1, 3.0, 0.85, "MP-01 Patrimonial TI", "#C5DDD4")
    _box(ax, 3.6, 4.1, 3.0, 0.85, "MP-02 Mantenimiento", "#C5DDD4")
    _box(ax, 6.9, 4.1, 3.0, 0.85, "MP-03 Gobernanza", "#C5DDD4")
    _box(ax, 0.3, 2.5, 1.45, 0.9, "P-Inventario", "#D6E4F0", fontsize=7)
    _box(ax, 1.85, 2.5, 1.45, 0.9, "P-Ficha", "#D6E4F0", fontsize=7)
    _box(ax, 3.6, 2.5, 1.45, 0.9, "P-Intervención", "#D6E4F0", fontsize=7)
    _box(ax, 5.15, 2.5, 1.45, 0.9, "P-Prioriz. ML", "#D6E4F0", fontsize=7)
    _box(ax, 6.9, 2.5, 1.45, 0.9, "P-Auth/RBAC", "#D6E4F0", fontsize=7)
    _box(ax, 8.45, 2.5, 1.45, 0.9, "P-Config.", "#D6E4F0", fontsize=7)
    _arrow(ax, 1.8, 4.1, 1.05, 3.4)
    _arrow(ax, 1.8, 4.1, 2.55, 3.4)
    _arrow(ax, 5.1, 4.1, 4.3, 3.4)
    _arrow(ax, 5.1, 4.1, 5.85, 3.4)
    _arrow(ax, 8.4, 4.1, 7.6, 3.4)
    _arrow(ax, 8.4, 4.1, 9.15, 3.4)
    _box(ax, 0.15, 0.35, 1.7, 1.3, "PR-Registrar\nequipo", "#E8EEF6", fontsize=7)
    _box(ax, 1.95, 0.35, 1.7, 1.3, "PR-Carga\nExcel", "#E8EEF6", fontsize=7)
    _box(ax, 4.25, 0.35, 1.9, 1.3, "PR-Registrar\nintervención", "#E8EEF6", fontsize=7)
    _arrow(ax, 1.05, 2.5, 1.0, 1.7)
    _arrow(ax, 1.05, 2.5, 2.8, 1.7)
    _arrow(ax, 4.3, 2.5, 5.2, 1.7)
    fig.savefig(IMG / "fig_a7_3_jerarquia.png", bbox_inches="tight", facecolor="white")
    plt.close()

    for name, title in (
        ("fig_a7_4_bizagi_asis.png", "PEGAR AQUÍ el export BIZAGI\nBPMN as-is (gestión actual)"),
        ("fig_a7_5_bizagi_tobe.png", "PEGAR AQUÍ el export BIZAGI\nBPMN to-be (flujo Sigemad)"),
    ):
        fig, ax = _new_fig(10.2, 4.0)
        ax.set_xlim(0, 10.2)
        ax.set_ylim(0, 4.0)
        ax.add_patch(Rectangle((0.4, 0.4), 9.4, 3.2, fill=False, ls="--", lw=1.6, edgecolor="#1A5F7A"))
        ax.text(5.1, 2.0, title, ha="center", va="center", fontsize=12, color="#1A5F7A")
        fig.savefig(IMG / name, bbox_inches="tight", facecolor="white")
        plt.close()

    fig, ax = _new_fig(10.2, 4.6)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 4.6)
    ax.text(0.3, 4.2, "prompts/", fontsize=11, fontweight="bold", color="#1F3A5F", family="monospace")
    lines = [
        "01_requisitos/           R-001 … R-005",
        "02_diseno/               D-001 … D-006",
        "03_implementacion/       I-001 … I-009  (+ 2 Superados)",
        "04_testing/              T-001 … T-005",
        "05_mantenimiento/        M-001 … M-004",
        "gobernanza/              plantilla · catálogo · política · métricas",
    ]
    for i, line in enumerate(lines):
        y = 3.5 - i * 0.52
        _box(ax, 0.4, y - 0.15, 9.4, 0.45, "", "#F4F7FB", fontsize=1)
        ax.text(0.6, y + 0.05, line, fontsize=8.5, family="monospace", color="#1A1A1A", va="center")
    fig.savefig(IMG / "fig_a7_6_arbol.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 2.8)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 2.8)
    fases = [
        (0.2, "R\nR-001…005", "#C5DDD4"),
        (2.2, "D\nD-001…006", "#D6E4F0"),
        (4.2, "I\nI-001…009", "#B8CDE0"),
        (6.2, "T\nT-001…005", "#F3E4C8"),
        (8.2, "M\nM-001…004", "#E8EEF6"),
    ]
    for x, t, c in fases:
        _box(ax, x, 0.85, 1.7, 1.25, t, c)
    for x in (1.9, 3.9, 5.9, 7.9):
        _arrow(ax, x, 1.45, x + 0.3, 1.45)
    fig.savefig(IMG / "fig_a7_7_ciclo.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 2.6)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 2.6)
    vers = ["0.1.0\nI-001", "0.2.0\nI-002", "0.3.0\nI-003", "0.4.0\nI-004", "0.5.0\nI-005",
            "0.6.0\nI-006", "0.7.0\nI-007", "0.8–0.9\nI-008", "0.9.1\nI-009"]
    for i, t in enumerate(vers):
        x = 0.15 + i * 1.12
        _box(ax, x, 0.7, 1.02, 1.2, t, "#D6E4F0", fontsize=6.6)
        if i < 8:
            _arrow(ax, x + 1.02, 1.3, x + 1.12, 1.3)
    fig.savefig(IMG / "fig_a7_8_incrementos.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 3.2)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 3.2)
    steps = [(0.3, "Medición"), (2.5, "Diagnóstico"), (4.7, "Refinamiento"), (6.9, "Decisión")]
    for x, t in steps:
        _box(ax, x, 1.5, 1.9, 1.0, t, "#B8CDE0")
    for x in (2.2, 4.4, 6.6):
        _arrow(ax, x, 2.0, x + 0.3, 2.0)
    _box(ax, 6.9, 0.2, 1.9, 0.85, "Aprobado → cierre", "#C5DDD4", fontsize=7)
    _arrow(ax, 7.85, 1.5, 7.85, 1.05)
    fig.savefig(IMG / "fig_a7_9_d2.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 2.6)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 2.6)
    _box(ax, 0.3, 0.7, 3.0, 1.2, "I-00N.md\nprompts/03_implementacion", "#D6E4F0", fontsize=8)
    _box(ax, 3.7, 0.7, 2.8, 1.2, "Commit GitHub\nej. f086a63", "#B8CDE0", fontsize=8)
    _box(ax, 6.9, 0.7, 3.0, 1.2, "Archivo en ese SHA\nEquipoController.php", "#C5DDD4", fontsize=8)
    _arrow(ax, 3.3, 1.3, 3.7, 1.3)
    _arrow(ax, 6.5, 1.3, 6.9, 1.3)
    fig.savefig(IMG / "fig_a7_11_cadena.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = _new_fig(10.2, 3.4)
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 3.4)
    levels = [
        (0.3, "N1", "Manual /\ncopia ocasional", "#E8EEF6", False),
        (2.25, "N2", "Asistido\n(este caso)", "#7EB0C8", True),
        (4.2, "N3", "Más autonomía\nen un repo", "#E8EEF6", False),
        (6.15, "N4", "Agentes sin\nrevisión por paso", "#E8EEF6", False),
        (8.1, "N5", "Autonomía\nen producción", "#E8EEF6", False),
    ]
    for x, n, s, c, mark in levels:
        _box(ax, x, 1.1, 1.8, 1.6, f"{n}\n{s}", c, fontsize=7.5)
        if mark:
            ax.text(x + 0.9, 2.95, "▲ alcanzado", ha="center", fontsize=8, color="#1A5F7A", fontweight="bold")
    fig.savefig(IMG / "fig_a7_12_niveles.png", bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = plt.subplots(figsize=(8.4, 3.8), dpi=160)
    fig.patch.set_facecolor("white")
    fases = ["Requisitos", "Diseño", "Implementación", "Pruebas", "Mantenimiento"]
    vals = [5, 6, 9, 5, 4]
    bars = ax.bar(fases, vals, color="#1A5F7A", width=0.55)
    ax.set_ylabel("Prompts vigentes")
    ax.set_ylim(0, 11)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.25, str(v), ha="center", fontsize=9)
    fig.savefig(IMG / "fig_a7_14_barras.png", bbox_inches="tight", facecolor="white")
    plt.close()


def build_doc():
    draw_all_figures()
    doc = Document()
    configure_section(doc.sections[0])

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cover, before=12, after=4, line=1.15, first_line=0)
    r = cover.add_run("UNIVERSIDAD CONTINENTAL")
    _set_run_font(r, size=12, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(sub, before=0, after=16, line=1.15, first_line=0)
    r = sub.add_run("Escuela Académico Profesional de Ingeniería de Sistemas e Informática")
    _set_run_font(r, size=11, color=GRAY)

    add_heading_custom(doc, "ANEXO 07. Desarrollo metodológico", 0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4, line=1.15, first_line=0)
    r = p.add_run("Plantilla guía — Intervención metodológica")
    _set_run_font(r, size=14, bold=True, color=TEAL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=12, line=1.15, first_line=0)
    r = p.add_run("Aplicación del Prompt-Centered SDLC v1.2  ·  Caso Sigemad MPA V2\nMunicipalidad Provincial de Acobamba")
    _set_run_font(r, size=11, italic=True, color=GRAY)

    add_body(
        doc,
        "Este archivo es una plantilla para guiar la redacción del Anexo 07. "
        "No es el anexo final de la tesis. Las cajas grises indican qué redactar; "
        "las azules, qué citar. Las tablas y figuras son estructuras de ejemplo "
        "(datos borrador de Sigemad). Sustituya [completar], pegue sus BPMN de Bizagi "
        "y borre las notas de uso antes de entregar.",
    )
    add_callout(
        doc,
        "Citar la fuente primaria de la metodología (Rojas Camayo, Prompt-Centered SDLC versión 1.2) "
        "al presentar el anexo. Este anexo alimenta el § 4.1 del Capítulo IV; no incluye "
        "conclusiones generales ni contrastación de hipótesis.",
        fill="D6E8F5",
        lead="Cita. ",
    )

    add_heading_custom(doc, "Nota de propósito y vínculo con la tesis", 1)
    add_callout(
        doc,
        "Indicar que este Anexo 07 documenta la intervención metodológica (aplicación del "
        "Prompt-Centered SDLC en el desarrollo de Sigemad MPA para la Municipalidad Provincial "
        "de Acobamba) y que la evidencia (prompts versionados, D2, matriz de doble entrada, "
        "trazabilidad y métricas de proceso) alimenta el Capítulo 4. Aclarar que las "
        "conclusiones y recomendaciones finales de la tesis no se desarrollan aquí.",
    )

    add_heading_custom(doc, "Contenido del anexo", 1)
    for item in [
        "1. Introducción a la intervención metodológica",
        "2. Justificación de la elección del Prompt-Centered SDLC",
        "3. Descripción del proyecto de software y de los procesos institucionales",
        "4. Preparación y entrada mínima del proyecto",
        "5. Aplicación por fases del SDLC",
        "6. Evaluación, refinamiento y versionado de prompts (D2)",
        "7. Matriz de doble entrada: trazabilidad de funciones y prompts",
        "8. Nivel de automatización agentica alcanzado",
        "9. Seguridad, gobernanza y monitoreo",
        "10. Trazabilidad de artefactos (síntesis)",
        "11. Métricas de proceso de la intervención",
        "12. Lecciones aprendidas de la intervención",
        "13. Cierre del anexo",
        "14. Referencias utilizadas en este anexo",
        "Apéndices A–G",
    ]:
        add_body(doc, item)

    # --- 1 ---
    add_heading_custom(doc, "1. Introducción a la intervención metodológica", 1)
    add_heading_custom(doc, "1.1 Propósito de este anexo", 2)
    add_callout(
        doc,
        "Redactar qué se documenta, qué metodología y versión se aplicó, sobre qué sistema, "
        "y qué se busca evidenciar (aplicación completa, trazabilidad, soporte al Capítulo 4).",
    )
    add_callout(doc, "Fuente primaria Prompt-Centered SDLC (Rojas Camayo, versión 1.2).", fill="D6E8F5", lead="Cita. ")
    add_body(doc, "Use la Figura 1 para explicar en una vista la unidad de trabajo de la intervención.")
    add_picture_centered(doc, IMG / "fig_a7_1_cadena.png", 15.5)
    add_caption(doc, "Figura 1. Cadena de la intervención: prompt, artefacto, métrica y decisión.")
    add_note(doc, "Elabore el párrafo citando esta figura. Si refina un prompt, el ciclo vuelve al primer recuadro.")

    add_heading_custom(doc, "1.2 Objetivos de la intervención", 2)
    add_heading_custom(doc, "1.2.1 Objetivo general", 3)
    add_callout(
        doc,
        "Enunciar el objetivo general de ESTA intervención (aplicar la metodología de requisitos "
        "a despliegue con evidencia), distinto del objetivo general de la tesis (efecto del sistema / ML).",
    )
    add_heading_custom(doc, "1.2.2 Objetivos específicos", 3)
    add_callout(
        doc,
        "Listar 4 a 6 objetivos verificables: anatomía D1, versionado, D2, matriz, automatización, "
        "guardrails, trazabilidad extremo a extremo.",
    )
    add_table(
        doc,
        ["N.º", "Objetivo específico de la intervención", "Evidencia", "Apartado"],
        [
            ["OE-1", "Aplicar anatomía D1 en todas las fases del SDLC", "29 registros vigentes", "4 y 5"],
            ["OE-2", "Evaluar salidas con ciclo D2 y dejar decisión", "registro_metricas.md", "6"],
            ["OE-3", "Trazar cada incremento a un prompt I-*", "I-001 … I-009; 0.1.0–0.9.1", "5.3"],
            ["OE-4", "Construir la matriz función–prompt–commit", "F-001 … F-008 + Excel V3", "7"],
            ["OE-5", "Aplicar guardrails y política de uso de IA", "politicas_uso_ia.md", "9"],
            ["OE-6", "Dejar trazabilidad HU → prompt → código → test", "Tabla 19", "10"],
        ],
        [Cm(1.4), Cm(6.4), Cm(4.5), Cm(3.2)],
    )
    add_caption(doc, "Tabla 1. Objetivos específicos de la intervención y dónde se evidencian.")
    add_note(doc, "Ajuste los OE si su asesor pide otro recorte. No copie aquí los objetivos del Capítulo I.")

    add_heading_custom(doc, "1.3 Alcance", 2)
    add_callout(
        doc,
        "Definir qué está incluido (fases SDLC, LLM/herramienta, revisión humana, registro de prompts, "
        "matriz) y qué queda fuera (modelos propios, N4–N5, módulos no prioritarios, ML obligatorio en producción).",
    )

    add_heading_custom(doc, "1.4 Marco metodológico de referencia", 2)
    add_callout(
        doc,
        "Presentar las cinco dimensiones D1–D5 y qué evidencia generó cada una. Mencionar la matriz "
        "de doble entrada (D3 / mantenimiento). Citar Rojas Camayo; SWEBOK solo si se usa.",
        fill="D6E8F5",
        lead="Cita. ",
    )
    add_picture_centered(doc, IMG / "fig_a7_2_d1d5.png", 15.5)
    add_caption(doc, "Figura 2. Dimensiones D1–D5 y evidencia generada en el caso Sigemad MPA.")

    # --- 2 ---
    add_heading_custom(doc, "2. Justificación de la elección del Prompt-Centered SDLC", 1)
    add_heading_custom(doc, "2.1 Problema metodológico", 2)
    add_callout(
        doc,
        "Describir el uso artesanal de IA (sin traza, sin criterios, sin gobernanza) y por qué en tesis "
        "y entidad pública se requiere un proceso auditable.",
    )
    add_callout(
        doc,
        "Uno o dos artículos sobre problemas del uso no estructurado de GenAI en desarrollo de software "
        "(p. ej. Paladino y Pons, 2025; Shrivastava et al., 2025).",
        fill="D6E8F5",
        lead="Cita. ",
    )
    add_table(
        doc,
        ["Limitación de Scrum con IA sistemática", "Qué se usó en Sigemad"],
        [
            ["El sprint dimensiona capacidad humana de implementación", "Incrementos por módulo (0.1.0–0.9.1), sin caja de dos semanas"],
            ["Ceremonias asumen coordinación de implementadores", "El humano aprueba, rechaza o pide refinamiento"],
            ["El backlog de sprint es la traza principal", "Traza: prompt → artefacto → métrica → decisión"],
            ["Historias estimables en un sprint", "Historias como técnica de requisitos; cierre con DoD + D2"],
        ],
        [Cm(7.6), Cm(7.9)],
    )
    add_caption(doc, "Tabla 2. Scrum frente a Prompt-Centered SDLC en este caso.")

    add_heading_custom(doc, "2.2 Posicionamiento frente a estándares", 2)
    add_heading_custom(doc, "2.2.1 ISO/IEC/IEEE 12207", 3)
    add_callout(
        doc,
        "La norma es el marco de procesos del ciclo de vida. Mapear las fases usadas con los procesos "
        "técnicos. La metodología se superpone a este marco; no lo sustituye. Cita obligatoria de la 12207.",
        fill="D6E8F5",
        lead="Cita. ",
    )
    add_heading_custom(doc, "2.2.2 ISO/IEC 29110", 3)
    add_callout(doc, "Equipo pequeño / VSE (un tesista). Citar el perfil Basic o el overview de la 29110.", fill="D6E8F5", lead="Cita. ")
    add_heading_custom(doc, "2.2.3 Otros cuerpos de conocimiento", 3)
    add_callout(doc, "SWEBOK, CMMI o el libro de la universidad, solo si se mencionan.", fill="D6E8F5", lead="Cita. ")
    add_table(
        doc,
        ["Fase", "Prompts", "Proceso técnico 12207 (verificar edición)", "Producto en Sigemad"],
        [
            ["Requisitos", "R-001 … R-005", "Stakeholder / System requirements definition", "Ficha, 56 HU, personas, ambigüedades, RNF"],
            ["Diseño", "D-001 … D-006", "Architecture / Design definition", "ER, ADR-001, ADR-002, contrato API"],
            ["Implementación", "I-001 … I-009", "Implementation", "Incrementos 0.1.0–0.9.1"],
            ["Pruebas", "T-001 … T-005", "Verification / Validation", "Plan, PHPUnit, Vitest, pytest"],
            ["Mantenimiento", "M-001 … M-004", "Transition / Maintenance", "producción, impacto, rollback"],
        ],
        [Cm(3.0), Cm(3.0), Cm(5.2), Cm(4.3)],
        font_size=8,
    )
    add_caption(doc, "Tabla 3. Fases del Prompt-Centered SDLC y procesos técnicos de ISO/IEC/IEEE 12207.")
    add_note(doc, "Ajuste los nombres de proceso a la edición de la norma que cite (p. ej. 2017).")

    add_heading_custom(doc, "2.3 Posicionamiento frente a la literatura de IA en el SDLC", 2)
    add_callout(
        doc,
        "Para cada artículo: autores, año, qué propone, fases que cubre, similitudes y diferencias. "
        "Síntesis: dejan una brecha (el prompt no es artefacto central; falta D2 y gobernanza). "
        "El Prompt-Centered la cierra sobre 12207/29110.",
    )
    add_table(
        doc,
        ["Fuente", "Qué aporta", "Qué no cubre", "Brecha que cierra esta propuesta"],
        [
            ["Yas, Alazzawi y Rahmatullah (2023)", "Fases estables del SDLC", "El prompt como artefacto", "Se conservan R/D/I/T/M; la IA entra dentro de cada fase"],
            ["[completar: Shrivastava / Hymel / Paladino]", "[completar]", "[completar]", "[completar]"],
            ["[completar tercera fuente]", "[completar]", "[completar]", "[completar]"],
        ],
        [Cm(4.0), Cm(3.7), Cm(3.6), Cm(4.2)],
        font_size=8,
    )
    add_caption(doc, "Tabla 4. Literatura de IA en el SDLC: aporta, no cubre y brecha.")

    add_heading_custom(doc, "2.4 Aporte diferencial del Prompt-Centered SDLC", 2)
    add_callout(
        doc,
        "Listar: prompt como artefacto central, trazabilidad, D2, seguridad desde el inicio, "
        "automatización gradual, matriz de doble entrada, documentabilidad. Citar Rojas Camayo.",
        fill="D6E8F5",
        lead="Cita. ",
    )
    add_heading_custom(doc, "2.5 Alineación con los objetivos de la tesis", 2)
    add_callout(
        doc,
        "Explicar cómo esta intervención aporta evidencia a la validación de la metodología y al § 4.1. "
        "No contrastar hipótesis aquí.",
    )

    # --- 3 ---
    add_heading_custom(doc, "3. Descripción del proyecto de software y de los procesos institucionales", 1)
    add_heading_custom(doc, "3.1 Ficha del proyecto", 2)
    add_callout(
        doc,
        "Completar ficha: nombre, entidad, objetivo, usuarios, problema, alcance incluido/excluido, "
        "stack, equipo, duración, nivel de automatización objetivo.",
    )
    add_table(
        doc,
        ["Campo", "Contenido (borrador — revíselo)"],
        [
            ["Nombre", "Sigemad MPA — Gestión de equipos y mantenimiento predictivo"],
            ["Código", "SIGEMAD-MPA-V2"],
            ["Entidad", "Municipalidad Provincial de Acobamba (informática / patrimonio TI)"],
            ["Versión de producto", "0.9.1"],
            ["Objetivo", "Inventario, fichas, historial de mantenimiento y ML opcional de riesgo"],
            ["Usuarios", "Administrador, Técnico, Practicante"],
            ["Stack", "React 19 + Vite, PHP 8.1+ REST, MySQL, FastAPI/Scikit-learn (opcional)"],
            ["Fuera de alcance", "Portal ciudadano, bienes no TI, SIGA/SIAF, ML obligatorio en producción"],
            ["Metodología de construcción", "Prompt-Centered SDLC v1.2"],
            ["Herramienta de IA", "Cursor Agent"],
            ["Nivel de automatización", "N2 (asistido; revisión humana obligatoria)"],
        ],
        [Cm(4.8), Cm(10.7)],
    )
    add_caption(doc, "Tabla 5. Ficha del proyecto Sigemad MPA V2.")

    add_heading_custom(doc, "3.2 Contexto y restricciones", 2)
    add_callout(doc, "Restricciones técnicas, de seguridad/privacidad, de tiempo, de infraestructura y de capacidad del equipo.")
    add_heading_custom(doc, "3.3 Justificación del caso", 2)
    add_callout(doc, "Por qué Acobamba permite demostrar la metodología hasta despliegue y con exigencia de trazabilidad.")

    add_heading_custom(doc, "3.4 Identificación de procesos (macroprocesos y procedimientos)", 2)
    add_callout(
        doc,
        "Pegue aquí SU inventario ya elaborado. Una fila = un procedimiento. "
        "La última columna: Sí / Parcial / No (automatizado en Sigemad). Las filas de ejemplo son modelo, no el catálogo final.",
    )
    add_table(
        doc,
        ["Código", "Macroproceso", "Procedimiento", "Responsable", "¿En Sigemad?"],
        [
            ["MP-01 / PR-01", "Gestión patrimonial TI", "Registrar equipo (código 12 dígitos)", "Técnico / Admin", "Sí — Inventario"],
            ["MP-01 / PR-02", "Gestión patrimonial TI", "Carga masiva por Excel", "Administrador", "Sí — I-009"],
            ["MP-02 / PR-01", "Mantenimiento de equipos", "Registrar ficha de intervención", "Técnico", "Sí — Mantenimiento"],
            ["MP-02 / PR-02", "Mantenimiento de equipos", "Consultar riesgo ML", "Técnico / Admin", "Parcial — FastAPI"],
            ["[completar]", "[completar con su catálogo]", "[completar]", "[completar]", "Sí / Parcial / No"],
        ],
        [Cm(2.6), Cm(3.6), Cm(4.4), Cm(2.5), Cm(2.4)],
        font_size=8,
    )
    add_caption(doc, "Tabla 6. Identificación de procesos (estructura de ejemplo; rellenar con su inventario).")
    add_picture_centered(doc, IMG / "fig_a7_3_jerarquia.png", 15.5)
    add_caption(doc, "Figura 3. Jerarquía de procesos (ejemplo). Sustituya los nodos por sus códigos reales.")

    add_heading_custom(doc, "3.5 Diagrama de procesos (BIZAGI)", 2)
    add_callout(
        doc,
        "Exporte PNG desde Bizagi (no capture la ventana). En el cuerpo: as-is, to-be y, si cabe, un procedimiento. "
        "El resto al Apéndice G.",
    )
    add_picture_centered(doc, IMG / "fig_a7_4_bizagi_asis.png", 15.5)
    add_caption(doc, "Figura 4. Diagrama de proceso as-is. Fuente: elaboración propia en Bizagi. [PEGAR EXPORT]")
    add_picture_centered(doc, IMG / "fig_a7_5_bizagi_tobe.png", 15.5)
    add_caption(doc, "Figura 5. Diagrama de proceso to-be. Fuente: elaboración propia en Bizagi. [PEGAR EXPORT]")

    # --- 4 ---
    add_heading_custom(doc, "4. Preparación y entrada mínima del proyecto", 1)
    add_heading_custom(doc, "4.1 Política mínima de uso de IA", 2)
    add_callout(
        doc,
        "Transcribir o adaptar: revisión humana obligatoria, no pegar datos sensibles, versionado de prompts, "
        "ADR, verificación con pruebas, capacitación. Fuente: prompts/gobernanza/politicas_uso_ia.md.",
    )
    add_heading_custom(doc, "4.2 Repositorio de prompts", 2)
    add_callout(doc, "Describir la estructura real de carpetas y la convención de nombres.")
    add_picture_centered(doc, IMG / "fig_a7_6_arbol.png", 15.5)
    add_caption(doc, "Figura 6. Árbol del repositorio de prompts del caso Sigemad MPA.")

    add_heading_custom(doc, "4.3 Definition of Done ampliada", 2)
    add_callout(doc, "Listar condiciones para cerrar un incremento cuando se usa IA.")
    add_table(
        doc,
        ["N.º", "Criterio", "¿Se exige?", "Evidencia"],
        [
            ["1", "Criterios de aceptación verificables", "Sí", "HU + incremento"],
            ["2", "Prompt (si hubo IA) versionado en prompts/", "Sí", "Catálogo D3"],
            ["3", "Revisión humana de lo integrado", "Sí", "Campo Revisor"],
            ["4", "Pruebas relevantes en verde", "Sí", "T-002 / T-003 / T-004"],
            ["5", "Sin secretos en el repositorio", "Sí", "Política + M-004"],
            ["6", "Documentación de la fase actualizada", "Sí", "documents/0N_*/"],
            ["7", "ADR si cambió la arquitectura", "Sí", "ADR-001, ADR-002"],
            ["8", "Trazabilidad HU → incremento → versión → módulo", "Sí", "Matriz + § 10"],
        ],
        [Cm(1.3), Cm(7.2), Cm(2.2), Cm(4.8)],
    )
    add_caption(doc, "Tabla 7. Definition of Done ampliada para trabajo asistido por IA.")

    add_heading_custom(doc, "4.4 Herramientas y modelos de IA", 2)
    add_table(
        doc,
        ["Herramienta o modelo", "Uso principal", "Fase(s)"],
        [
            ["Cursor Agent", "Generación y refinamiento de artefactos", "R, D, I, T, M"],
            ["[completar: modelo del registro]", "[completar]", "[completar]"],
            ["GitHub", "Versionado de código y SHA de evidencia", "I, T, M"],
            ["PHPUnit / Vitest / pytest", "Ejecución de pruebas (el humano interpreta)", "T"],
            ["Bizagi", "Modelado BPMN de procesos institucionales", "Entrada de R (§ 3)"],
        ],
        [Cm(5.2), Cm(6.6), Cm(3.7)],
    )
    add_caption(doc, "Tabla 8. Herramientas y modelos utilizados por fase.")

    # --- 5 ---
    add_heading_custom(doc, "5. Aplicación por fases del SDLC", 1)
    add_callout(
        doc,
        "En cada fase: objetivo, entradas, técnicas, prompts (código + versión + modelo; texto largo en apéndice), "
        "outputs, D2, decisión y entregables. No recitar la 12207 en cada subfase.",
    )
    add_picture_centered(doc, IMG / "fig_a7_7_ciclo.png", 15.5)
    add_caption(doc, "Figura 7. Ciclo R–D–I–T–M y familia de prompts vigentes.")
    add_picture_centered(doc, IMG / "fig_a7_8_incrementos.png", 15.5)
    add_caption(doc, "Figura 8. Línea de incrementos de producto (no sprints Scrum).")

    add_heading_custom(doc, "5.1 Fase de requisitos", 2)
    add_callout(doc, "Objetivo, entradas (incluye el mapa de procesos del § 3.4–3.5), técnicas, R-01/R-02 y resto, D2, entregables.")
    add_table(
        doc,
        ["Código", "Técnica", "Output / artefacto", "Estado", "Decisión D2"],
        [
            ["R-001", "[completar]", "ficha-proyecto.md", "Reconstruido", "Aprobado"],
            ["R-002", "[completar]", "56 HU / 8 épicas", "Reconstruido", "Aprobado"],
            ["R-003", "[completar]", "personas.md", "Reconstruido", "Aprobado"],
            ["R-004", "[completar]", "ambiguedades.md", "Ejecutado", "Aprobado"],
            ["R-005", "[completar]", "requisitos_no_funcionales.md", "Ejecutado", "Aprobado"],
        ],
        [Cm(2.2), Cm(2.8), Cm(5.3), Cm(2.8), Cm(2.4)],
        font_size=8,
    )
    add_caption(doc, "Tabla 9. Fase de requisitos: prompt, output y decisión D2.")

    add_heading_custom(doc, "5.2 Fase de diseño", 2)
    add_callout(doc, "Objetivo, entradas, técnicas, prompts de modelo de datos y arquitectura, ADR, D2, entregables.")
    add_table(
        doc,
        ["Código", "Técnica", "Output", "ADR", "Estado", "Decisión D2"],
        [
            ["D-001", "[completar]", "v2_estructura.sql", "—", "Reconstruido", "Aprobado"],
            ["D-001 v2", "[completar]", "er_v2.md", "—", "Ejecutado", "Aprobado"],
            ["D-002", "[completar]", "Stack", "ADR-001", "Reconstruido", "Aprobado"],
            ["D-003", "[completar]", "contrato_api_v2.md", "—", "Ejecutado", "Aprobado"],
            ["D-004", "[completar]", "frontend_features.md", "—", "Ejecutado", "Aprobado"],
            ["D-005", "[completar]", "Análisis ML", "—", "Reconstruido", "Aprobado"],
            ["D-006", "[completar]", "Degradación FastAPI", "ADR-002", "Ejecutado", "Aprobado"],
        ],
        [Cm(2.2), Cm(2.4), Cm(3.6), Cm(2.2), Cm(2.6), Cm(2.5)],
        font_size=8,
    )
    add_caption(doc, "Tabla 10. Fase de diseño: prompt, ADR y decisión D2.")

    add_heading_custom(doc, "5.3 Fase de implementación", 2)
    add_callout(doc, "Documentar 2–3 prompts representativos en el párrafo; la tabla cubre I-001…I-009 (un prompt ≈ un incremento).")
    add_table(
        doc,
        ["Código", "Incremento / versión", "Output principal", "Estado", "Decisión D2"],
        [
            ["I-001", "1 / 0.1.0", "Cimientos, API, DDL", "Reconstruido", "Aprobado"],
            ["I-002", "2 / 0.2.0", "CRUD inventario", "Reconstruido", "Aprobado"],
            ["I-003", "3 / 0.3.0", "Fichas y mantenimiento", "Reconstruido", "Aprobado"],
            ["I-004", "4 / 0.4.0", "Reportes PDF", "Reconstruido", "Aprobado"],
            ["I-005", "5 / 0.5.0", "Configuración", "Reconstruido", "Aprobado"],
            ["I-006", "— / 0.6.0", "Auth JWT + dashboard", "Reconstruido", "Aprobado"],
            ["I-007", "6 / 0.7.0", "Microservicio ML", "Reconstruido", "Aprobado"],
            ["I-008", "7 / 0.8.0–0.9.0", "Telemetría y ficha predictiva", "Reconstruido", "Aprobado"],
            ["I-009", "Parche / 0.9.1", "RBAC + plantilla Excel", "Reconstruido", "Aprobado"],
        ],
        [Cm(2.0), Cm(3.4), Cm(4.7), Cm(2.8), Cm(2.6)],
        font_size=8,
    )
    add_caption(doc, "Tabla 11. Fase de implementación: un prompt por incremento.")
    add_table(
        doc,
        ["Función", "Módulo", "Prompt", "Commit", "Archivo ancla"],
        [
            ["F-001", "auth", "I-006", "5ce7575", "src/features/auth/Login.jsx"],
            ["F-002", "configuracion", "I-005", "4e08b9e", "ConfiguracionPage.jsx"],
            ["F-003", "inventario", "I-002", "f086a63", "EquipoController.php"],
            ["F-004", "ficha", "I-003", "f086a63", "FichaTecnicaController.php"],
            ["F-005", "mantenimiento", "I-003", "8b20337", "MantenimientoController.php"],
            ["F-006", "dashboard", "I-006", "5ce7575", "DashboardPage.jsx"],
            ["F-007", "ml", "I-007", "e9a0965", "ml/app/main.py"],
            ["F-008", "reportes", "I-004", "f086a63", "ReporteController.php"],
        ],
        [Cm(2.0), Cm(3.0), Cm(2.2), Cm(2.6), Cm(5.7)],
        font_size=8,
    )
    add_caption(doc, "Tabla 14. Incremento, SHA y archivo ancla (puente con la matriz del § 7).")
    add_note(doc, "I-008 e I-009 son evoluciones del mismo módulo. No borre el SHA de origen. No duplique esta tabla en el § 7.")

    add_heading_custom(doc, "5.4 Fase de testing", 2)
    add_table(
        doc,
        ["Código", "Runner o entregable", "Output", "Estado", "Decisión D2"],
        [
            ["T-001", "Plan funcional (humano)", "plan_pruebas_funcionales.md", "Reconstruido", "Aprobado"],
            ["T-002", "PHPUnit", "backend/tests/", "Reconstruido", "Aprobado"],
            ["T-003", "Vitest", "src/lib/*.test.js", "Reconstruido", "Aprobado"],
            ["T-004", "pytest", "ml/tests/", "Reconstruido", "Aprobado"],
            ["T-005", "Plantilla de diagnóstico", "diagnostico_test_fallido.md", "Plantilla", "Aprobado"],
        ],
        [Cm(2.0), Cm(4.2), Cm(4.3), Cm(2.6), Cm(2.4)],
        font_size=8,
    )
    add_caption(doc, "Tabla 12. Fase de pruebas.")

    add_heading_custom(doc, "5.5 Fase de mantenimiento y despliegue", 2)
    add_callout(doc, "Incluir plan de despliegue, rollback, pruebas de humo y responsable. Entregables: sistema, docs, prompts.")
    add_table(
        doc,
        ["Código", "Output", "Estado", "Decisión D2"],
        [
            ["M-001", "produccion.md (despliegue)", "Reconstruido", "Aprobado"],
            ["M-002", "Plantilla de análisis de impacto", "Plantilla", "Aprobado"],
            ["M-003", "Plantilla de documentación post-cambio", "Plantilla", "Aprobado"],
            ["M-004", "Rollback y secretos en produccion.md", "Ejecutado", "Aprobado"],
        ],
        [Cm(2.4), Cm(7.5), Cm(3.0), Cm(2.6)],
    )
    add_caption(doc, "Tabla 13. Fase de mantenimiento y despliegue.")

    # --- 6 ---
    add_heading_custom(doc, "6. Evaluación, refinamiento y versionado de prompts (D2)", 1)
    add_heading_custom(doc, "6.1 Ciclo D2 aplicado", 2)
    add_callout(doc, "Describir medición → diagnóstico → refinamiento/versionado. Citar Rojas Camayo al explicar D2.", fill="D6E8F5", lead="Cita. ")
    add_picture_centered(doc, IMG / "fig_a7_9_d2.png", 15.5)
    add_caption(doc, "Figura 9. Ciclo D2: medición, diagnóstico, refinamiento y decisión.")

    add_heading_custom(doc, "6.2 Registro de versiones", 2)
    add_table(
        doc,
        ["Código", "Fase", "Iteraciones", "Medición D2", "Decisión", "Lección breve"],
        [
            ["R-004", "Requisitos", "[completar]", "14 hallazgos R-01", "Aprobado", "Estado honesto (cerrado/supuesto)"],
            ["D-006", "Diseño", "1", "ADR-002 opción D", "Aprobado", "Lo tácito no existe para la tesis"],
            ["I-002", "Implementación", "No medido", "F-003 / HU-INV-001–002", "Aprobado", "I-002 ya no es ML"],
            ["I-007", "Implementación", "≥2", "Acc. 95 % / F1 0.92 (modelo, no hipótesis)", "Aprobado", "Schema {} no []"],
            ["[completar]", "[completar]", "[completar]", "desde registro_metricas.md", "[completar]", "[completar]"],
        ],
        [Cm(1.8), Cm(2.6), Cm(2.2), Cm(3.6), Cm(2.0), Cm(3.3)],
        font_size=8,
    )
    add_caption(doc, "Tabla 15. Extracto del registro de versiones (D2).")

    add_heading_custom(doc, "6.3 Ejemplo de refinamiento", 2)
    add_callout(doc, "Un solo caso real: original, resultado, evidencia, diagnóstico, cambios, criterio de cierre, decisión.")
    add_table(
        doc,
        ["Campo", "Antes", "Después"],
        [
            ["Identificador", "I-001-MACRO (api + auth + inventario)", "I-001 … I-009 (un prompt por incremento)"],
            ["Problema", "Un prompt cubría varios módulos", "Ambigüedad de ID (I-002 inventario vs I-002-ML)"],
            ["Decisión D2", "Superado", "Aprobado (vigentes)"],
            ["Lección", "Un I-* = un incremento", "Los macros no se borran; quedan Superados"],
        ],
        [Cm(3.2), Cm(6.15), Cm(6.15)],
        font_size=8,
    )
    add_caption(doc, "Tabla 10b. Ejemplo de refinamiento: de macro-prompt a I-001…I-009.")

    # --- 7 ---
    add_heading_custom(doc, "7. Matriz de doble entrada: trazabilidad de funciones y prompts", 1)
    add_heading_custom(doc, "7.1 Propósito de la matriz", 2)
    add_callout(
        doc,
        "Instrumento que conecta, por función: ciclo de vida y arquitectura; capacidades DevOps; "
        "prompt de origen; ruta/commit; test asociado. Sirve para mantenimiento y para evidenciar la metodología.",
    )
    add_picture_centered(doc, IMG / "fig_a7_11_cadena.png", 15.5)
    add_caption(doc, "Figura 11. Cadena de evidencia: prompt versionado → commit GitHub → archivo en ese SHA.")

    add_heading_custom(doc, "7.2 Fundamento teórico de la matriz", 2)
    add_callout(
        doc,
        "Amaro et al. (capacidades DevOps y LCPs); Jiang et al. (versionado y calidad de prompts). "
        "La operacionalización conjunta y el enlace al repositorio es aporte propio.",
        fill="D6E8F5",
        lead="Cita. ",
    )
    add_heading_custom(doc, "7.3 Estructura de la matriz", 2)
    add_callout(doc, "Describir bloques de columnas: función, impacto LCP, cinco capacidades DevOps, datos del prompt, repo, test.")
    add_heading_custom(doc, "7.4 Aplicación de la matriz en este caso", 2)
    add_callout(doc, "Alcance F-001…F-008 + I-008/I-009. Hallazgos de madurez DevOps. Referenciar el Excel (Apéndice A).")
    add_table(
        doc,
        ["Función", "Prompt", "Commit", "Control de versiones", "CI", "Tests", "Monitoreo"],
        [
            ["F-003", "I-002", "f086a63", "Sí (GitHub)", "[completar]", "[completar]", "[completar]"],
            ["F-007", "I-007", "e9a0965", "Sí", "[completar]", "pytest", "[completar]"],
            ["…", "usar Tabla 14", "—", "Sí/Parcial/No", "Sí/Parcial/No", "Sí/Parcial/No", "Sí/Parcial/No"],
        ],
        [Cm(2.0), Cm(2.0), Cm(2.2), Cm(2.8), Cm(2.1), Cm(2.2), Cm(2.2)],
        font_size=8,
    )
    add_caption(doc, "Tabla 16. Extracto DevOps de la matriz (no duplique la Tabla 14).")
    add_heading_custom(doc, "7.5 Uso de la matriz para mantenimiento", 2)
    add_callout(doc, "Procedimiento ante un cambio futuro: ubicar función → DevOps/LCP → prompt → código → test → actualizar versión.")
    add_heading_custom(doc, "7.6 Evidencia adjunta", 2)
    add_callout(doc, "La matriz completa (Excel) es el Apéndice A. Evidencia primaria función–prompt–código–test.")

    # --- 8 ---
    add_heading_custom(doc, "8. Nivel de automatización agentica alcanzado", 1)
    add_heading_custom(doc, "8.1 Nivel alcanzado", 2)
    add_callout(doc, "Indicar N1–N5 y qué significó. Opcional: citar un artículo de agentic AI para contrastar.", fill="D6E8F5", lead="Cita. ")
    add_picture_centered(doc, IMG / "fig_a7_12_niveles.png", 15.5)
    add_caption(doc, "Figura 12. Escala N1–N5. En este caso se alcanzó N2 (asistido / agentico acotado).")
    add_heading_custom(doc, "8.2 Justificación de no avanzar a niveles superiores", 2)
    add_callout(doc, "Entidad pública, secretos, tesis, necesidad de revisión humana. No N4–N5.")
    add_heading_custom(doc, "8.3 Reglas de control", 2)
    add_table(
        doc,
        ["Regla", "Qué prohíbe o exige", "Dónde está escrita"],
        [
            ["Revisión humana antes de integrar", "Copiar código a ciegas", "Política de IA"],
            ["No secretos en prompt ni repo", "Credenciales, tokens", "D1 «No hacer»; M-004"],
            ["Un I-* por incremento", "Macros que mezclan módulos", "Oleada 2"],
            ["Criterio de parada", "Iterar sin decisión D2", "Aprobado / Rechazado / Dividir"],
            ["No N4–N5", "Agentes sin supervisión en producción", "§ 8.2"],
        ],
        [Cm(5.0), Cm(5.5), Cm(5.0)],
        font_size=8,
    )
    add_caption(doc, "Tabla 17. Reglas de control del nivel de automatización.")

    # --- 9 ---
    add_heading_custom(doc, "9. Seguridad, gobernanza y monitoreo", 1)
    add_heading_custom(doc, "9.1 Guardrails en los prompts", 2)
    add_callout(doc, "Tabla riesgo / guardrail / control. Opcional: papers de riesgos GenAI o secure SDLC.", fill="D6E8F5", lead="Cita. ")
    add_table(
        doc,
        ["Riesgo", "Guardrail en el prompt (D1)", "Control posterior"],
        [
            ["Alucinación de APIs o tablas", "Restricciones + entradas (SQL/ADR existentes)", "Revisión humana + pruebas"],
            ["Secretos en el chat o en el repo", "«No secretos» en No hacer", "local.php fuera de git; M-004"],
            ["Alcance inflado (módulos inventados)", "Tarea acotada al incremento", "DoD + un I-* por incremento"],
            ["Código inseguro (UI sin RBAC)", "Criterios de aceptación", "I-009 requireRole; tests"],
            ["Reescritura total / IP", "No reescribir V1; no dependencias no autorizadas", "Diff de GitHub"],
        ],
        [Cm(4.6), Cm(5.6), Cm(5.3)],
        font_size=8,
    )
    add_caption(doc, "Tabla 18. Riesgo, guardrail en el prompt y control posterior.")
    add_heading_custom(doc, "9.2 Monitoreo", 2)
    add_callout(doc, "Dimensiones, indicadores, frecuencia y acción si falla.")
    add_heading_custom(doc, "9.3 Cumplimiento de la política de uso de IA", 2)
    add_callout(doc, "Confirmar el § 4.1. Ser honesto si la auditoría periódica está pendiente.")

    # --- 10 ---
    add_heading_custom(doc, "10. Trazabilidad de artefactos (síntesis)", 1)
    add_callout(
        doc,
        "Matriz corta HU → prompt → diseño → código → test → versión. El detalle fino está en el § 7. "
        "Tres o cuatro filas bastan.",
    )
    add_table(
        doc,
        ["HU (ejemplo)", "Prompt", "Diseño", "Código", "Test", "Versión"],
        [
            ["HU-INV-001 / 002", "I-002", "Contrato /equipos; ER v2_equipos", "EquipoController.php", "[completar]", "0.2.0"],
            ["[completar HU-MANT]", "I-003", "Ficha de mantenimiento", "MantenimientoController.php", "[completar]", "0.3.0"],
            ["[completar HU-ML]", "I-007", "D-005, ADR-002", "ml/app/main.py", "pytest T-004", "0.7.0"],
            ["AUTH-005 / CFG-006", "I-009", "RBAC", "requireRole + Navbar", "T-002", "0.9.1"],
        ],
        [Cm(2.8), Cm(1.8), Cm(4.2), Cm(3.4), Cm(1.8), Cm(1.5)],
        font_size=8,
    )
    add_caption(doc, "Tabla 19. Síntesis de trazabilidad HU → prompt → diseño → código → test → versión.")

    # --- 11 ---
    add_heading_custom(doc, "11. Métricas de proceso de la intervención", 1)
    add_callout(
        doc,
        "Solo métricas de PROCESO (tiempos, n.º de prompts, iteraciones, tasa de aprobación, DoD, DevOps). "
        "La estadística inferencial y las hipótesis van al Capítulo 4.",
    )
    add_table(
        doc,
        ["Indicador", "Valor", "Criterio", "Estado"],
        [
            ["Historias de usuario implementadas", "56", "Alcance cerrado", "OK"],
            ["Incrementos de implementación", "7 + parche 0.9.1", "Trazables a I-*", "OK"],
            ["Prompts vigentes", "29 (R×5, D×6, I×9, T×5, M×4)", "Mapa D3 por fase", "OK"],
            ["Prompts superados", "2", "No borrar", "OK"],
            ["Anatomía D1 completa", "29/29", "Plantilla maestra", "OK"],
            ["Ciclo D2 documentado", "Sí; iteraciones I-* no medidas en caliente", "Decisión + lección", "Parcial"],
            ["ADR", "2", "Decisiones de arquitectura", "OK"],
            ["Reconstrucción vs ejecución", "Mixto", "Registrar antes el próximo cambio", "Parcial"],
        ],
        [Cm(4.6), Cm(5.0), Cm(3.8), Cm(2.1)],
        font_size=8,
    )
    add_caption(doc, "Tabla 20. Tablero de métricas de proceso (no inferencial).")
    add_picture_centered(doc, IMG / "fig_a7_14_barras.png", 14.0)
    add_caption(doc, "Figura 14. Prompts vigentes por fase (no incluye los 2 Superados).")

    # --- 12 ---
    add_heading_custom(doc, "12. Lecciones aprendidas de la intervención", 1)
    add_callout(doc, "Lecciones técnicas y metodológicas, recomendaciones prácticas para otra aplicación, limitaciones del caso.")
    add_table(
        doc,
        ["Hallazgo", "Mejora abierta"],
        [
            ["Parte de los I-* se documentó a posteriori", "Próximo cambio de código: registrar el prompt antes"],
            ["I-001 fue un macro-prompt; iteraciones no medidas en caliente", "Oleada 2 ya partió I-001…I-009; medir en caliente de aquí en adelante"],
            ["Un solo producto municipal no generaliza", "Comparar con otro caso (p. ej. SGMI) o con Scrum clásico"],
            ["ML opcional; FastAPI no es obligatorio en producción", "El caso no cubre un SDLC de ML en producción"],
            ["Nombres residuales de la etapa Scrum (verify_sprint6_ml.py)", "Incremento de higiene, sin reescribir historia"],
        ],
        [Cm(7.6), Cm(7.9)],
        font_size=8,
    )
    add_caption(doc, "Tabla 21. Hallazgos de la intervención y mejoras abiertas.")

    # --- 13 ---
    add_heading_custom(doc, "13. Cierre del anexo", 1)
    add_callout(
        doc,
        "Cierre breve: se aplicó la metodología de forma documentada; hay evidencia de prompts, D2, "
        "trazabilidad (matriz) y despliegue; esta evidencia soporta el Capítulo 4. "
        "No incluir conclusiones ni recomendaciones generales de la tesis. Puede citarse de nuevo Rojas Camayo.",
        fill="D6E8F5",
        lead="Cita. ",
    )

    # --- 14 ---
    add_heading_custom(doc, "14. Referencias utilizadas en este anexo", 1)
    add_callout(doc, "Solo fuentes citadas en ESTE anexo, estilo Continental. Deben coincidir con las cajas azules.")
    add_heading_custom(doc, "14.1 Estándares y normas", 2)
    add_body(doc, "ISO/IEC/IEEE 12207; ISO/IEC 29110; SWEBOK (si se citó).")
    add_heading_custom(doc, "14.2 Metodología propia", 2)
    add_body(doc, "Rojas Camayo. Prompt-Centered SDLC versión 1.2. [completar ficha bibliográfica].")
    add_heading_custom(doc, "14.3 Artículos sobre IA en el SDLC", 2)
    add_body(doc, "Yas et al. (2023); [completar Shrivastava / Hymel / Paladino y Pons] — los de la Tabla 4.")
    add_heading_custom(doc, "14.4 Fuentes de la matriz de doble entrada", 2)
    add_body(doc, "Amaro, Pereira y Mira da Silva; Jiang et al. La operacionalización conjunta es aporte propio.")
    add_heading_custom(doc, "14.5 Libros y otras fuentes", 2)
    add_body(doc, "Huyen (2025); Krishna y Meda (2024); otros solo si se citaron en el anexo.")

    add_heading_custom(doc, "Apéndices del Anexo 07", 1)
    add_table(
        doc,
        ["Apéndice", "Contenido", "Estado"],
        [
            ["A", "Matriz de doble entrada (Excel V3 + prompts detallados)", "Existe en documents/matriz_doble_entrada/"],
            ["B", "Catálogo de prompts versionados (índice, no los 29 textos)", "prompts/gobernanza/catalogo.md"],
            ["C", "Outputs representativos (1 HU, 1 ADR, 1 fragmento, 1 test)", "[completar selección]"],
            ["D", "Checklists D1/D2 de 2–3 prompts", "[completar]"],
            ["E", "Reportes de cobertura / pruebas", "[completar si existen]"],
            ["F", "Evidencias de revisión y despliegue a producción", "[completar]"],
            ["G", "Macroprocesos, procedimientos y BPMN Bizagi", "[pegar sus archivos]"],
        ],
        [Cm(2.2), Cm(9.5), Cm(3.8)],
        font_size=8,
    )
    add_caption(doc, "Tabla 22. Apéndices previstos del Anexo 07.")

    add_body(
        doc,
        "Fin de la plantilla guía. Conservar Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC_solo_indicaciones.docx "
        "como la versión solo-texto. Al redactar el anexo definitivo, elimine las cajas grises/azules y deje "
        "prosa académica + tablas/figuras numeradas al estilo Continental.",
    )

    doc.save(OUT)
    print("OK", OUT)


if __name__ == "__main__":
    build_doc()

# -*- coding: utf-8 -*-
"""Redibuja las cinco dimensiones D1–D5 y las ancla en el apartado 1.4."""
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent
DOC = BASE / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
IMG = BASE / "_figuras_anexo7" / "fig_a7_2_d1d5.png"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)


def _font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_body(p, size=12, first_indent=0.75, after=8, justify=True, center=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    )
    pf = p.paragraph_format
    pf.space_before = Pt(0 if not center else 4)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0 if center else 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0 if center else first_indent)


def set_runs_text(p, text, size=12, bold=False, italic=False, justify=True, first_indent=0.75, after=8, center=False):
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size, bold=bold, italic=italic)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size, bold=bold, italic=italic)
    style_body(p, size=size, first_indent=first_indent, after=after, justify=justify, center=center)


def insert_after(element, doc):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    return Paragraph(new_p, doc)


def insert_body_after(element, text, doc):
    p = insert_after(element, doc)
    run = p.add_run(text)
    _font(run, 12)
    style_body(p)
    return p


def has_text(doc, needle):
    return any(needle in p.text for p in doc.paragraphs)


def _round(ax, x, y, w, h, facecolor, edge="#1F3A5F", lw=1.25):
    ax.add_patch(
        FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            linewidth=lw, edgecolor=edge, facecolor=facecolor,
        )
    )


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1), (x2, y2),
            arrowstyle="-|>", mutation_scale=11, linewidth=1.2, color="#1F3A5F",
        )
    )


def draw_d1d5():
    IMG.parent.mkdir(exist_ok=True)
    plt.rcParams["font.family"] = "Times New Roman"
    w, h = 11.4, 7.15
    fig, ax = plt.subplots(figsize=(w, h), dpi=170)
    fig.patch.set_facecolor("white")
    ax.axis("off")
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)

    ax.text(
        5.7, 6.90, "Cinco dimensiones del Prompt-Centered SDLC",
        ha="center", va="center", fontsize=14, fontweight="bold", color="#1F3A5F",
    )
    ax.text(
        5.7, 6.58, "D4 gobierna el uso de IA; D1–D3 operan el prompt por fase; D5 deja la traza auditable.",
        ha="center", va="center", fontsize=8.4, color="#4A5568",
    )

    # D4 roof
    _round(ax, 1.55, 5.15, 8.30, 1.22, "#E8D5C4")
    ax.text(5.70, 6.05, "D4  ·  Gobernanza", ha="center", va="center",
            fontsize=11.5, fontweight="bold", color="#1F3A5F")
    ax.text(5.70, 5.62, "¿Con qué reglas se usa el agente?   Política mínima de IA, nivel N2, revisión humana.",
            ha="center", va="center", fontsize=8.0, color="#1A1A1A")
    ax.text(5.70, 5.32, "Evidencia en Sigemad: politicas_uso_ia.md  ·  Tabla 6",
            ha="center", va="center", fontsize=7.4, color="#4A5568")

    _arrow(ax, 5.70, 5.15, 5.70, 4.88)

    # D1 D2 D3
    mid = [
        (0.28, "D1  ·  Anatomía", "¿Cómo se formula el pedido?",
         "Rol, tarea, restricciones\ny bloque «No hacer».",
         "Plantilla y registros\nR / D / I / T / M", "#C5DDD4"),
        (4.02, "D2  ·  Evaluación", "¿Se aprueba la salida?",
         "Medir, diagnosticar,\nrefinar y decidir.",
         "registro_metricas.md\ny bloque D2 del prompt", "#D6E4F0"),
        (7.76, "D3  ·  Mapa por fase", "¿En qué fase del SDLC?",
         "Un código por fase:\nR, D, I, T o M.",
         "Catálogo de prompts\nvigentes", "#B8CDE0"),
    ]
    bw, bh, y = 3.36, 2.05, 2.72
    for x, title, q, how, ev, col in mid:
        _round(ax, x, y, bw, bh, col)
        ax.text(x + bw / 2, y + 1.76, title, ha="center", va="center",
                fontsize=10.2, fontweight="bold", color="#1F3A5F")
        ax.text(x + bw / 2, y + 1.32, q, ha="center", va="center",
                fontsize=8.0, color="#1A1A1A")
        ax.text(x + bw / 2, y + 0.78, how, ha="center", va="center",
                fontsize=7.5, color="#4A5568", linespacing=1.15)
        ax.text(x + bw / 2, y + 0.28, ev, ha="center", va="center",
                fontsize=7.3, color="#1A1A1A", linespacing=1.12)
    _arrow(ax, 3.64, y + bh / 2, 4.02, y + bh / 2)
    _arrow(ax, 7.38, y + bh / 2, 7.76, y + bh / 2)

    _arrow(ax, 5.70, 2.72, 5.70, 2.42)

    # D5 floor
    _round(ax, 1.55, 0.98, 8.30, 1.32, "#F3E4C8")
    ax.text(5.70, 1.98, "D5  ·  Trazabilidad", ha="center", va="center",
            fontsize=11.5, fontweight="bold", color="#1F3A5F")
    ax.text(5.70, 1.55, "¿Se puede auditar sin abrir el chat?   Función → prompt → commit GitHub → archivo.",
            ha="center", va="center", fontsize=8.0, color="#1A1A1A")
    ax.text(5.70, 1.20, "Evidencia en Sigemad: matriz de doble entrada V3  ·  Tabla 14  ·  Figura 13",
            ha="center", va="center", fontsize=7.4, color="#4A5568")

    ax.add_patch(Rectangle((0.22, 0.14), 10.96, 0.70, facecolor="#F4F7FB", edgecolor="#1F3A5F", lw=1.0))
    ax.text(
        5.70, 0.49,
        "Lectura: D4 pone el techo (qué se puede pedirle a la IA). D1 escribe el prompt, D2 lo cierra y D3 lo sitúa en una fase. D5 es el piso: la evidencia pública.",
        ha="center", va="center", fontsize=7.6, color="#1A1A1A",
    )

    fig.savefig(IMG, bbox_inches="tight", facecolor="white", pad_inches=0.08)
    plt.close()
    return IMG


INTRO = (
    "El marco se opera en cinco dimensiones, ilustradas en la Figura 2. D4 (gobernanza) es el techo: "
    "la política mínima de uso de IA, el nivel N2 y la revisión humana obligatoria. Debajo operan tres "
    "dimensiones de trabajo: D1 (anatomía) formula el pedido al modelo, con restricciones y bloque "
    "«No hacer»; D2 (evaluación) mide, diagnostica, refina y deja una decisión escrita; D3 (mapa por fase) "
    "sitúa cada prompt en requisitos, diseño, implementación, pruebas o mantenimiento. D5 (trazabilidad) "
    "es el piso: función, prompt, commit y archivo, sin depender del historial de chat. Cada dimensión "
    "se evidencia más adelante en el apartado que le corresponde."
)

AFTER = (
    "La Figura 2 se lee de arriba abajo. Arriba, D4 responde con qué reglas se usa el agente. En el centro, "
    "D1, D2 y D3 forman el ciclo del prompt dentro de una fase del SDLC. Abajo, D5 permite que un tercero "
    "abra el prompt, el commit y el archivo en GitHub. La política de IA no es un texto suelto: es D4 "
    "operando sobre las otras cuatro dimensiones."
)


def replace_picture_before_caption(doc, caption_start, image_path, width_cm=16.0):
    paras = list(doc.paragraphs)
    last_i = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith(caption_start):
            last_i = i
    if last_i is None:
        return False
    for j in range(last_i - 1, max(-1, last_i - 6), -1):
        q = paras[j]
        drawings = q._element.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
        )
        if drawings or not q.text.strip():
            for child in list(q._element):
                tag = child.tag
                if tag.endswith("}r") or tag.endswith("}hyperlink"):
                    q._element.remove(child)
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.first_line_indent = Cm(0)
            q.paragraph_format.space_before = Pt(6)
            q.paragraph_format.space_after = Pt(2)
            run = q.add_run()
            run.add_picture(str(image_path), width=Cm(width_cm))
            return True
    return False


def main():
    img = draw_d1d5()
    print("PNG", img)
    doc = Document(str(DOC))

    heading = None
    for p in doc.paragraphs:
        if p.text.strip() == "1.4 Marco metodológico de referencia":
            heading = p
    if heading is None:
        raise RuntimeError("No se encontró 1.4")

    seen = False
    body = None
    for p in doc.paragraphs:
        if p is heading:
            seen = True
            continue
        if not seen:
            continue
        t = p.text.strip()
        if not t:
            continue
        if t.startswith(("Tabla ", "Figura ", "2.")):
            break
        body = p
        break
    if body is not None:
        set_runs_text(body, INTRO)

    caps = [p for p in doc.paragraphs if p.text.strip().startswith("Figura 2.")]
    if caps:
        set_runs_text(
            caps[-1],
            "Figura 2. Cinco dimensiones del Prompt-Centered SDLC: gobernanza (D4), anatomía, "
            "evaluación y mapa por fase (D1–D3) y trazabilidad (D5).",
            size=11, italic=True, center=True, first_indent=0, after=8,
        )
        if not has_text(doc, "La Figura 2 se lee de arriba abajo"):
            insert_body_after(caps[-1]._element, AFTER, doc)

    ok = replace_picture_before_caption(doc, "Figura 2.", img, 16.0)
    print("Figura reemplazada", ok)
    doc.save(str(DOC))
    print("OK", DOC)


if __name__ == "__main__":
    main()

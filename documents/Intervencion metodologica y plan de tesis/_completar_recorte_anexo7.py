# -*- coding: utf-8 -*-
"""Segunda pasada: solo residuales. No vuelve a insertar [17]–[21] ni toca 3.4–3.5."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

DOC = Path(__file__).resolve().parent / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

_BLOCK35_ELS = set()


def _font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_paragraph(p, size=12, first_indent=0.75, after=8, justify=True):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_indent:
        pf.first_line_indent = Cm(first_indent)
    for run in p.runs:
        _font(run, size=size)


def insert_paragraph_after(element, text, size=12, first_indent=0.75):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    p = Paragraph(new_p, element.getparent())
    run = p.add_run(text)
    _font(run, size=size)
    style_paragraph(p, size=size, first_indent=first_indent)
    return p


def set_cell(cell, text, size=9):
    p = cell.paragraphs[0]
    for extra in list(cell.paragraphs)[1:]:
        extra._element.getparent().remove(extra._element)
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size)


def compute_block_35(doc):
    global _BLOCK35_ELS
    kids = list(doc.element.body)
    start = end = None
    for i, el in enumerate(kids):
        if not el.tag.endswith("}p"):
            continue
        t = "".join(el.itertext()).strip()
        if t.startswith("3.4 Identificación") or t.startswith("3.4 Identificacion"):
            start = i
        if start is not None and (t.startswith("4. Preparación") or t.startswith("4. Preparacion")):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError("No se localizó el bloque 3.4–4")
    _BLOCK35_ELS = set(kids[start:end])
    print(f"Bloque 3.4–3.5 protegido: {end - start} elementos.")


def in_35_para(p):
    return p._element in _BLOCK35_ELS


def replace_para_text(p, new_text, size=12, first_indent=0.75):
    if p.runs:
        p.runs[0].text = new_text
        _font(p.runs[0], size=size)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(new_text)
        _font(run, size=size)
    style_paragraph(p, size=size, first_indent=first_indent)


def main():
    doc = Document(str(DOC))
    compute_block_35(doc)

    # 1) Figure 1 de Scrum (no tocar Figura 1 de la cadena)
    for p in doc.paragraphs:
        if in_35_para(p):
            continue
        t = p.text.replace("\xa0", " ").strip()
        if t.startswith("Figure 1") and "Scrum" in t:
            replace_para_text(
                p,
                "La comparación operativa se resume en la Tabla 2.",
                first_indent=0.75,
            )
            print("Figure 1 Scrum -> remision a Tabla 2")

    # 2) 2.4: cinco viñetas
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("2.4 Aporte"):
            nxt = doc.paragraphs[i + 1]
            if nxt.text.strip().startswith("El aporte de esta intervención se resume"):
                print("2.4 ya convertido")
                break
            bullets = [
                "El prompt se trata como artefacto de ingeniería y se versiona.",
                "Cada incremento se cierra con una decisión D2.",
                "La unidad de planificación es el incremento acotado, no el sprint.",
                "La matriz de doble entrada mantiene la traza prompt–commit–archivo.",
                "Las restricciones de seguridad se escriben desde la anatomía del prompt (D1).",
            ]
            replace_para_text(nxt, "El aporte de esta intervención se resume en lo siguiente:", first_indent=0.75)
            anchor = nxt._element
            for b in reversed(bullets):
                insert_paragraph_after(anchor, "• " + b, first_indent=0.5)
            print("2.4 convertido a cinco vinetas")
            break

    # 3) 5.4: párrafo breve (faltaba)
    for p in doc.paragraphs:
        if p.text.strip() == "5.4 Fase de testing":
            # no duplicar si ya hay prosa
            break
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "5.4 Fase de testing":
            nxt = doc.paragraphs[i + 1].text.strip()
            if nxt.startswith("Tabla 12"):
                insert_paragraph_after(
                    p._element,
                    "En pruebas se ejecutaron el plan funcional y los runners asociados a T-001 "
                    "a T-005. El responsable interpretó los resultados; T-005 permanece como "
                    "plantilla de diagnóstico de fallos. El registro y la decisión D2 se "
                    "resumen en la Tabla 12.",
                )
                print("5.4: párrafo de fase añadido")
            break

    # 4) 5.5: quitar “tesista”
    for p in doc.paragraphs:
        if in_35_para(p):
            continue
        if "El tesista actuó" in p.text:
            replace_para_text(
                p,
                p.text.replace(
                    "El tesista actuó como desarrollador y revisor.",
                    "El autor actuó como desarrollador y revisor.",
                ),
            )
            print("5.5: tesista -> autor")

    # 5) Caption Tabla 16 (tono de plantilla)
    for p in doc.paragraphs:
        if in_35_para(p):
            continue
        if "no duplique la Tabla 14" in p.text:
            replace_para_text(
                p,
                "Tabla 16. Extracto DevOps de la matriz (los identificadores SHA constan en la Tabla 14).",
                first_indent=0,
            )
            print("Caption Tabla 16")

    # 6) Celdas residuales
    for table in doc.tables:
        if table._tbl in _BLOCK35_ELS:
            continue
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text
                if txt.strip() == "Hostinger, impacto, rollback" or txt.strip() == "el ambiente de producción, impacto, rollback":
                    set_cell(cell, "Puesta en producción, impacto, rollback", 8)
                    print("Tabla 3: Hostinger -> produccion")
                    continue
                if "Contenido (borrador" in txt:
                    set_cell(cell, "Contenido", 9)
                    print("Tabla 5: encabezado sin 'borrador'")
                    continue
                if "hostinger.md" in txt or "Hostinger" in txt:
                    set_cell(
                        cell,
                        txt.replace("hostinger.md", "el plan de puesta en producción")
                        .replace("Hostinger", "el ambiente de producción"),
                        8,
                    )
                    print("Celda residual Hostinger:", txt[:60])

    # 7) 4.4: una línea si solo hay caption
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "4.4 Herramientas y modelos de IA":
            nxt = doc.paragraphs[i + 1].text.strip()
            if nxt.startswith("Tabla 8"):
                insert_paragraph_after(
                    p._element,
                    "Las herramientas y el destino de operación se recogen en la Tabla 8. "
                    "El entorno de edición y el agente se cuentan una sola vez; Git y GitHub "
                    "forman una fila; el destino es el ambiente de producción institucional.",
                )
                print("4.4: párrafo breve")
            break

    # 8) Nota Tabla 14 si falta
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Tabla 14."):
            nxt = doc.paragraphs[i + 1].text.strip() if i + 1 < len(doc.paragraphs) else ""
            if "apartado 7" not in nxt and "SHA" not in nxt:
                insert_paragraph_after(
                    p._element,
                    "Las evoluciones I-008 e I-009 no sustituyen el commit de origen. "
                    "El apartado 7 remite a esta tabla y no reproduce los identificadores SHA.",
                )
                print("Nota Tabla 14 añadida")
            break

    doc.save(str(DOC))
    print("OK", DOC)


if __name__ == "__main__":
    main()

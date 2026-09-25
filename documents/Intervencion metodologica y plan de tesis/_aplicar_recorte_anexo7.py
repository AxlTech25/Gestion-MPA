# -*- coding: utf-8 -*-
"""Recorte in-place del Anexo 07: menos redundancia, tono de informe, producción (no Hostinger).
No toca el bloque 3.4–3.5 (macroprocesos / Bizagi). No regenera el documento.
"""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

DOC = Path(__file__).resolve().parent / "Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC.docx"
FONT = "Times New Roman"
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x4A, 0x55, 0x68)


def _font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_paragraph(p, size=12, first_indent=0.75, after=8, justify=True):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_indent:
        pf.first_line_indent = Cm(first_indent)
    for run in p.runs:
        _font(run, size=size)


def insert_paragraph_after(element, text, size=12, first_indent=0.75):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    p = Paragraph(new_p, element.getparent())
    run = p.add_run(text)
    _font(run, size=size)
    style_paragraph(p, size=size, first_indent=first_indent)
    return p


def insert_paragraph_before(element, text, size=12, first_indent=0.75):
    new_p = OxmlElement("w:p")
    element.addprevious(new_p)
    p = Paragraph(new_p, element.getparent())
    run = p.add_run(text)
    _font(run, size=size)
    style_paragraph(p, size=size, first_indent=first_indent)
    return p


def delete_element(el):
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def is_1x1(table):
    return len(table.rows) == 1 and len(table.columns) == 1


def cell0(table):
    return table.cell(0, 0).text.strip()


def set_cell(cell, text, size=9):
    p = cell.paragraphs[0]
    for extra in list(cell.paragraphs)[1:]:
        delete_element(extra._element)
    if p.runs:
        p.runs[0].text = text
        _font(p.runs[0], size=size)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        _font(run, size=size)


# Elementos del body entre 3.4 y 4. Preparación (se calcula una vez).
_BLOCK35_ELS = set()


def _compute_block_35(doc):
    """Marca los hijos del body entre 3.4 y 4. Preparación para no tocarlos."""
    global _BLOCK35_ELS
    body = doc.element.body
    kids = list(body)
    start = end = None
    for i, el in enumerate(kids):
        if not el.tag.endswith("}p"):
            continue
        t = "".join(el.itertext()).strip()
        if t.startswith("3.4 Identificación") or t.startswith("3.4 Identificacion"):
            start = i
        if start is not None and (t.startswith("4. Preparación") or t.startswith("4. Preparacion")):
            end = i
            break
    if start is None or end is None:
        _BLOCK35_ELS = set()
        print("AVISO: no se localizó el bloque 3.4–4; no se aplicará la exclusión.")
        return
    _BLOCK35_ELS = set(kids[start:end])
    print(f"Bloque 3.4–3.5: {end - start} elementos (índices {start}–{end}).")


def _in_block_35(p):
    return p._element in _BLOCK35_ELS


def _table_in_block_35(table):
    return table._tbl in _BLOCK35_ELS


def replace_run_text(doc, mapping, skip_35=True):
    for p in doc.paragraphs:
        if skip_35 and _in_block_35(p):
            continue
        for run in p.runs:
            if not run.text:
                continue
            t = run.text
            for old, new in mapping:
                if old in t:
                    t = t.replace(old, new)
            run.text = t


def para_starts_with(p, *prefixes):
    t = p.text.strip()
    return any(t.startswith(pref) for pref in prefixes)


def find_para(doc, predicate):
    for p in doc.paragraphs:
        if predicate(p):
            return p
    return None


def convert_1x1_to_paragraph(table, rewrite=None):
    text = rewrite if rewrite is not None else cell0(table)
    tbl = table._tbl
    insert_paragraph_after(tbl, text)
    delete_element(tbl)


def delete_table(table):
    delete_element(table._tbl)


def delete_paragraph(p):
    delete_element(p._element)


def clear_table_keep_header(table, new_rows):
    """Replace body rows of a table with new_rows (list of list of str). Header kept."""
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        delete_element(tr)
    for vals in new_rows:
        row = table.add_row()
        for i, val in enumerate(vals):
            set_cell(row.cells[i], val, size=8)


def main():
    doc = Document(str(DOC))
    _compute_block_35(doc)

    # --- texto global (fuera de 3.4-3.5) ---
    text_map = [
        ("Plantilla guía — Intervención metodológica", "Informe de la intervención metodológica"),
        ("hostinger.md (despliegue)", "plan de puesta en producción"),
        ("hostinger.md", "plan de puesta en producción"),
        ("Hostinger (hPanel)", "Ambiente de producción institucional"),
        ("sin FastAPI en Hostinger", "sin servicio predictivo dedicado"),
        ("sin pipeline en Hostinger", "CI/CD no implantado en producción"),
        ("sin pipeline Hostinger", "sin pipeline de CI/CD en producción"),
        ("No (sin pipeline Hostinger)", "No (CI/CD no implantado en producción)"),
        ("Hostinger", "el ambiente de producción"),
        ("hPanel", "el servidor institucional"),
        ("ML obligatorio en Hostinger (el módulo se apaga si no hay FastAPI)",
         "ML obligatorio en producción (el módulo se desactiva si el servicio predictivo no está disponible)"),
        ("Checklist Hostinger + captura de publicación",
         "Checklist de puesta en producción, pruebas de humo y responsable"),
        ("Evidencias de revisión y despliegue Hostinger",
         "Evidencias de revisión y puesta en producción"),
        ("Figura 6. Árbol del repositorio de prompts del caso Sigemad MPA.",
         "Figura 9. Árbol del repositorio de prompts del caso Sigemad MPA."),
        ("Figura 7. Ciclo R–D–I–T–M y familia de prompts vigentes.",
         "Figura 10. Ciclo R–D–I–T–M y familia de prompts vigentes."),
        ("Figura 8. Línea de incrementos de producto (no sprints Scrum).",
         "Figura 11. Línea de incrementos de producto (no sprints Scrum)."),
        ("Figura 9. Ciclo D2: medición, diagnóstico, refinamiento y decisión.",
         "Figura 12. Ciclo D2: medición, diagnóstico, refinamiento y decisión."),
        ("Figura 11. Cadena de evidencia: prompt versionado → commit GitHub → archivo en ese SHA.",
         "Figura 13. Cadena de evidencia: prompt versionado, commit y archivo en ese SHA."),
        ("Figura 12. Escala N1–N5. En este caso se alcanzó N2 (asistido / agentico acotado).",
         "Figura 14. Escala N1–N5. En este caso se alcanzó el nivel N2 (asistido / agentico acotado)."),
        ("Figura 14. Prompts vigentes por fase (no incluye los 2 Superados).",
         "Figura 15. Prompts vigentes por fase (no incluye los dos prompts Superados)."),
        ("Tabla 3. Fases del Prompt-Centered SDLC y procesos técnicos de ISO/IEC/IEEE 12207.",
         "Tabla 3. Fases del Prompt-Centered SDLC y características de ISO/IEC 25010."),
        ("2.2.1 ISO/IEC/IEEE 12207", "2.2.1 ISO 9001"),
        ("2.2.2 ISO/IEC 29110", "2.2.2 ISO/IEC 25010"),
        ("2.2.3 ISO/IEC 9001", "2.2.3 ISO/IEC/IEEE 29119-2"),
        ("Figure 1 Comparación entre supuestos operativos de Scrum y de un SDLC centrado en prompts y agentes.",
         "La comparación operativa se resume en la Tabla 2."),
        ("Figure 1 Comparación entre supuestos operativos de Scrum y de un SDLC centrado en prompts y agentes.",
         "La comparación operativa se resume en la Tabla 2."),
    ]
    replace_run_text(doc, text_map, skip_35=True)

    # Párrafos de plantilla / duplicados de normas (fuera de 3.4-3.5)
    drop_exact = {
        "Borrador de trabajo del Anexo 07. Las cajas ya contienen prosa para revisar. Los apartados 3.4 y 3.5 (macroprocesos y Bizagi) se dejan para una pasada posterior. Elimine las cajas residuales cuando pase el texto al estilo Continental.",
        "Use la Figura 1 para explicar en una vista la unidad de trabajo de la intervención.",
        "La Figura 1 es la unidad de trabajo: si el prompt se refina, el ciclo vuelve al primer recuadro.",
        "Los OE de la Tabla 1 son de la intervención, no los del Capítulo I.",
        "La columna normativa de la Tabla 3 usa ISO/IEC 25010 (calidad de producto), no 12207.",
        "En este caso la calidad del proceso se ancla en ISO 9001 (registros, revisión y mejora = D1/D2). ISO/IEC 29110 queda como contexto VSE; la fase de pruebas se alineará después con ISO/IEC/IEEE 29119-2.",
        "ISO/IEC 25010 cubre las características de producto (mantenibilidad, eficiencia, seguridad) que la metodología ya opera como RNF y DoD; por eso no se usa ISO/IEC/IEEE 12207 como norma rectora.",
        "La Tabla 8 lista el stack real del caso (ADR-001) y las herramientas de IA, calidad y operación.",
        "I-008 e I-009 son evoluciones del mismo módulo. No borre el SHA de origen. No duplique esta tabla en el apartado 7.",
        "Fin de la plantilla guía. Conservar Anexo7_Intervencion_Metodologica_Prompt_Centered_SDLC_solo_indicaciones.docx como la versión solo-texto. Al redactar el anexo definitivo, elimine las cajas grises/azules y deje prosa académica + tablas/figuras numeradas al estilo Continental.",
        "14.1 Estándares y normas",
        "ISO 9001 (sistemas de gestión de la calidad). ISO/IEC 25010 (calidad del producto de software). ISO/IEC 29110 (contexto VSE). ISO/IEC/IEEE 29119-2 (procesos de prueba; alineación prevista). SWEBOK solo si se cita en el cuerpo.",
        "14.2 Metodología propia",
        "Rojas Camayo. Prompt-Centered SDLC versión 1.2. [completar ciudad, editorial o URL institucional].",
        "14.3 Artículos sobre IA en el SDLC",
        "Yas, Alazzawi y Rahmatullah (2023); Shrivastava et al. (2025); Hymel (AI-Native SDLC); Paladino y Pons (2025).",
        "14.4 Fuentes de la matriz de doble entrada",
        "Amaro, Pereira y Mira da Silva; Jiang et al. La operacionalización conjunta es aporte propio.",
        "14.5 Libros y otras fuentes",
        "Huyen (2025); Krishna y Meda (2024); otros solo si se citaron en el anexo.",
    }
    for p in list(doc.paragraphs):
        if _in_block_35(p):
            continue
        if p.text.strip() in drop_exact:
            delete_paragraph(p)

    # --- tablas: identificar por contenido ---
    tables = list(doc.tables)

    def find_tbl(pred):
        for t in tables:
            try:
                if pred(t):
                    return t
            except Exception:
                continue
        return None

    # 1.1: borrar T0-T3 y dejar un párrafo
    t0 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("Este anexo documenta la aplicación operativa"))
    t1 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("El Anexo 07 registra"))
    t2 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("Se documenta la aplicación completa"))
    t3 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("Fuente primaria: Rojas Camayo"))
    purpose = (
        "El presente anexo documenta la intervención metodológica con la que se construyó "
        "Sigemad MPA V2 en la Municipalidad Provincial de Acobamba, mediante la aplicación "
        "del Prompt-Centered SDLC versión 1.2. La unidad de trabajo se observa en la Figura 1: "
        "prompt versionado, artefacto, métrica y decisión. La evidencia de proceso alimenta el "
        "apartado 4.1 del Capítulo IV y no sustituye al Capítulo III ni a la contrastación de hipótesis."
    )
    if t0 is not None:
        insert_paragraph_before(t0._tbl, purpose)
    for t in (t0, t1, t2, t3):
        if t is not None:
            delete_table(t)

    # T5 caja OE (redundante con Tabla 1)
    t5 = find_tbl(lambda t: is_1x1(t) and "Los objetivos específicos de la intervención se resumen" in cell0(t))
    if t5 is not None:
        delete_table(t5)

    # T4 OG -> párrafo
    t4 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("Objetivo general de la intervención"))
    if t4 is not None:
        convert_1x1_to_paragraph(
            t4,
            "El objetivo general de esta intervención fue aplicar el Prompt-Centered SDLC v1.2, "
            "desde la elicitación de requisitos hasta la puesta en producción, dejando evidencia "
            "versionada de prompts, revisión humana y trazabilidad. Dicho objetivo no sustituye "
            "al objetivo general de la tesis.",
        )

    # Fusionar Scrum: reescribir T9 y borrar T10
    t9 = find_tbl(lambda t: (not is_1x1(t)) and t.cell(0, 0).text.strip() == "Aspecto"
                  and "Scrum" in t.cell(0, 1).text)
    t10 = find_tbl(lambda t: (not is_1x1(t)) and "Limitación de Scrum" in t.cell(0, 0).text)
    if t9 is not None:
        set_cell(t9.cell(0, 0), "Aspecto", 8)
        set_cell(t9.cell(0, 1), "Scrum clásico", 8)
        set_cell(t9.cell(0, 2), "Este caso (Prompt-Centered)", 8)
        merged = [
            ["Unidad de trabajo", "Sprints de 1 a 4 semanas",
             "Incrementos 0.1.0–0.9.1, sin caja fija de dos semanas"],
            ["Coordinación", "Equipo multifuncional y ceremonias",
             "Un desarrollador aprueba o rechaza la salida del agente"],
            ["Traza principal", "Backlog de sprint",
             "Prompt → artefacto → métrica → decisión"],
            ["Historias de usuario", "Ítems estimables en un sprint",
             "Técnica de requisitos; cierre con DoD y D2"],
            ["Riesgo dominante", "Comunicación y tamaño de equipo",
             "Auditabilidad, seguridad y calidad del prompt"],
        ]
        clear_table_keep_header(t9, merged)
    if t10 is not None:
        delete_table(t10)

    # T15 literatura prosa
    t15 = find_tbl(lambda t: is_1x1(t) and cell0(t).startswith("Yas et al. (2023) recuerdan"))
    if t15 is not None:
        delete_table(t15)

    # ISO boxes: reescribir y convertir a párrafo
    t11 = find_tbl(lambda t: is_1x1(t) and "12207" in cell0(t) and "ciclo de vida" in cell0(t))
    if t11 is not None:
        convert_1x1_to_paragraph(
            t11,
            "ISO 9001 se emplea para argumentar que el desarrollo asistido por IA debe dejar "
            "registros controlados: el prompt versionado, la revisión humana y la decisión D2 "
            "cumplen la función de planificación, verificación y acción correctiva. El ciclo de "
            "vida R/D/I/T/M ya está cubierto por la metodología aplicada; por ello no se adopta "
            "ISO/IEC/IEEE 12207 como norma rectora.",
        )
    t12 = find_tbl(lambda t: is_1x1(t) and "29110" in cell0(t) and "VSE" in cell0(t))
    if t12 is not None:
        convert_1x1_to_paragraph(
            t12,
            "ISO/IEC 25010 orienta las características de calidad del producto —en especial "
            "adecuación funcional, mantenibilidad, seguridad y eficiencia— ya operadas como "
            "requisitos no funcionales (R-005) y como Definition of Done. El mapeo por fase se "
            "presenta en la Tabla 3.",
        )
    t13 = find_tbl(lambda t: is_1x1(t) and "ISO 9001 se usa" in cell0(t))
    if t13 is not None:
        convert_1x1_to_paragraph(
            t13,
            "La fase de pruebas se alinea con los procesos de ISO/IEC/IEEE 29119-2 (plan, diseño "
            "e implementación de pruebas). ISO/IEC 29110 se menciona solo como contexto de un "
            "equipo muy pequeño (VSE), no como norma de diseño del ciclo.",
        )

    # Convertir el resto de cajas 1x1 a prosa, excepto 3.4/3.5
    skip_starts = (
        "Pendiente. Los apartados 3.4",
    )
    rewrites = {
        "Queda incluido:": (
            "El alcance de la intervención comprende las cinco fases del SDLC, el uso de un agente "
            "acotado con revisión humana, el repositorio de prompts vigentes, la matriz de doble "
            "entrada y la puesta en producción con componente de ML opcional. Quedan fuera el "
            "entrenamiento de un modelo fundacional propio, la automatización de niveles N4–N5, "
            "el portal ciudadano, el inventario no informático y la sustitución de SIGA/SIAF. "
            "Los macroprocesos y diagramas de proceso se desarrollan en los apartados 3.4 y 3.5."
        ),
        "El Prompt-Centered SDLC se opera": (
            "El marco se opera en cinco dimensiones, ilustradas en la Figura 2: D1, anatomía del "
            "prompt; D2, evaluación y decisión; D3, mapa de prompts por fase; D4, gobernanza; y D5, "
            "trazabilidad. Cada dimensión se evidencia más adelante en el apartado que le corresponde."
        ),
        "El aporte diferencial": (
            "El aporte de esta intervención se resume en cinco puntos: (1) el prompt se trata "
            "como artefacto de ingeniería versionado; (2) cada incremento se cierra con una "
            "decisión D2; (3) la unidad de planificación es el incremento acotado, no el sprint; "
            "(4) la matriz de doble entrada mantiene la traza prompt–commit–archivo; y (5) las "
            "restricciones de seguridad se escriben desde la anatomía del prompt (D1)."
        ),
        "Esta intervención no demuestra": (
            "La intervención no demuestra la hipótesis de la tesis. Aporta la evidencia de proceso "
            "que el apartado 4.1 debe resumir: un sistema real construido sin ceremonias Scrum, "
            "con incrementos acotados y con registro de lo ejecutado y de lo reconstruido a posteriori."
        ),
        "La identidad del caso": (
            "Sigemad MPA V2 es el software institucional del caso y, a la vez, el artefacto con el "
            "que se valida la metodología. Los datos de identificación se presentan en la Tabla 5."
        ),
        "Restricciones:": (
            "Las restricciones del caso fueron un solo desarrollador, datos patrimoniales y "
            "credenciales que no pueden incluirse en el prompt, un stack fijado por el ADR-001 y "
            "un ambiente de producción institucional. El componente de aprendizaje automático es "
            "opcional: si el servicio predictivo no está disponible, el resto del sistema opera "
            "con normalidad (ADR-002)."
        ),
        "Acobamba exige": (
            "El caso municipal exige inventario patrimonial, fichas técnicas, historial de "
            "mantenimiento y puesta en producción. Ello obliga a trazabilidad y a una política de "
            "secretos, y permite validar la metodología hasta el ambiente de operación, no solo "
            "hasta un prototipo de laboratorio."
        ),
        "Se aplicó la política v1.3": (
            "Se aplicó una política mínima de uso de IA: revisión humana antes de integrar; "
            "prohibición de pegar credenciales o datos personales en el prompt; versionado del "
            "prompt y no del chat; registro ADR ante cambios de arquitectura; verificación con "
            "pruebas; y un prompt de implementación por incremento. El detalle operativo se "
            "consolida en la Tabla 7 (Definition of Done)."
        ),
        "Los prompts viven": (
            "Los prompts se organizaron por fase en un repositorio versionado, según la Figura 9. "
            "La convención de confirmación en el control de versiones cita el código del prompt "
            "del incremento. No se versiona el historial de conversación."
        ),
        "Un incremento asistido": (
            "Un incremento asistido por IA no se consideró cerrado por el solo hecho de que el "
            "modelo hubiera generado código. Los criterios de cierre se recogen en la Tabla 7."
        ),
        "Cada fase se documenta": (
            "Cada fase se documentó con el mismo esquema: objetivo, entradas, técnica de prompting, "
            "salidas y decisión D2. El texto completo de los prompts se reserva al Apéndice B; "
            "en este apartado se presentan el código, la técnica y el artefacto."
        ),
        "Objetivo: cerrar alcance": (
            "En requisitos se cerraron el alcance, las personas, las historias, las ambigüedades "
            "y los requisitos no funcionales. Las técnicas predominantes fueron el encadenamiento "
            "de pensamiento guiado y el few-shot. Los registros y su decisión se resumen en la Tabla 9."
        ),
        "Objetivo: fijar datos": (
            "En diseño se fijaron el modelo de datos, el stack, el contrato de la API, el mapa "
            "del frontend y el recorte del componente predictivo. Las decisiones de arquitectura "
            "quedaron en los ADR-001 y ADR-002. El resumen aparece en la Tabla 10."
        ),
        "Objetivo: entregar software": (
            "En implementación se entregó el software por incrementos, con un prompt I-* por "
            "versión. Tres registros ilustran el patrón: I-002 (inventario y código patrimonial), "
            "I-007 (microservicio de ML) e I-009 (control de acceso tras un hallazgo de seguridad). "
            "Las Tablas 11 y 14 concentran incrementos y anclas en el repositorio."
        ),
        "El despliegue de producción": (
            "La puesta en producción se documentó con el registro M-001: secuencia de despliegue, "
            "pruebas de humo, responsable y plan de retorno. El rollback y la exclusión de secretos "
            "constan en M-004. M-002 y M-003 quedan como plantillas para el siguiente cambio de "
            "base de datos o de contrato. El tesista actuó como desarrollador y revisor."
        ),
        "El ciclo D2, según Rojas": (
            "El ciclo D2 consiste en medición, diagnóstico, refinamiento o versionado y decisión "
            "(aprobado, refinado, rechazado, dividir o superado). En este caso, las iteraciones "
            "de origen de varios registros I-* no se midieron en caliente, lo que se declara en "
            "la Tabla 15."
        ),
        "El caso de refinamiento": (
            "El refinamiento documentado es el paso del macro-prompt I-001 y del alias I-002-ML "
            "a la serie I-001 a I-009. Un identificador corresponde a un incremento; los macros "
            "se conservan con estado Superado, según la Tabla 10b."
        ),
        "La matriz de doble entrada V3 conecta": (
            "La matriz de doble entrada V3 conecta cada función del sistema con su prompt de origen, "
            "el commit correspondiente y el archivo en ese identificador. Su propósito es auditar "
            "la intervención y apoyar el mantenimiento. La cadena de evidencia se observa en la Figura 13; "
            "los SHA constan en la Tabla 14."
        ),
        "Se apoya en Amaro": (
            "El instrumento se apoya en Amaro y colaboradores, respecto de las capacidades DevOps "
            "y su relación con el ciclo de vida, y en Jiang y colaboradores, respecto del versionado "
            "y la calidad de prompts. La operacionalización conjunta y el enlace navegable al "
            "repositorio constituyen el aporte de este caso."
        ),
        "Columnas: identificación": (
            "La matriz agrupa identificación de la función, impacto sobre el ciclo de vida, "
            "cinco capacidades DevOps, datos del prompt, ubicación en el repositorio y prueba "
            "asociada. El detalle de columnas se adjunta en el Apéndice A."
        ),
        "Se aplicó a F-001": (
            "Se aplicó a las funciones F-001 a F-008 y a las evoluciones I-008 e I-009. El control "
            "de versiones está implantado. La integración y el despliegue continuos no se "
            "implantaron en el ambiente de producción de este caso. Las pruebas unitarias existen "
            "por entorno de ejecución. El extracto DevOps se limita a la Tabla 16; los SHA no se "
            "repiten (véase la Tabla 14)."
        ),
        "Ante un cambio futuro": (
            "Para un cambio futuro se propone: ubicar la función en la matriz; consultar el prompt "
            "de origen; abrir el archivo en el commit de ancla; ejecutar la prueba asociada; "
            "versionar un prompt nuevo si interviene IA; y actualizar iteraciones y estado DevOps."
        ),
        "La evidencia primaria es": (
            "La evidencia primaria de la matriz se adjunta como Apéndice A."
        ),
        "Se alcanzó el nivel N2": (
            "Se alcanzó el nivel N2, ilustrado en la Figura 14: un agente acotado genera artefactos "
            "y el responsable humano define el prompt, revisa, prueba e integra. No hubo "
            "encadenamiento autónomo de agentes en producción."
        ),
        "No se avanzó a N3": (
            "No se avanzó a los niveles N3 a N5 porque el caso corresponde a una entidad pública, "
            "existen secretos y datos patrimoniales, y la tesis exige poder explicar cada decisión."
        ),
        "Los guardrails se escriben": (
            "Las restricciones de seguridad se escribieron en la anatomía del prompt. El cruce "
            "entre riesgo, control en el prompt y control posterior se presenta en la Tabla 18. "
            "Las reglas generales de uso de IA constan en el apartado 4.1."
        ),
        "Se monitorea:": (
            "En el ambiente de producción se vigilan la disponibilidad de la API, la degradación "
            "del componente predictivo (ADR-002) y los incidentes de autenticación. No se implantó "
            "un monitor de aplicaciones de propósito general. La revisión se realiza en cada "
            "incremento y en cada publicación."
        ),
        "La política v1.3 se cumplió": (
            "La política de uso de IA se cumplió en lo esencial. Subsiste, como limitación, la "
            "ausencia de una auditoría periódica formal, señalada en el apartado 12."
        ),
        "La Tabla 19 resume": (
            "La Tabla 19 resume cuatro hilos de trazabilidad. El detalle por función se encuentra "
            "en el apartado 7 y en la Tabla 14."
        ),
        "Solo se reportan métricas": (
            "En este apartado solo se reportan métricas de proceso (Tabla 20 y Figura 15). La "
            "exactitud del modelo asociado a I-007 es una medición D2 de ese registro, no la "
            "contrastación de la hipótesis de la tesis."
        ),
        "Las lecciones (Tabla 21)": (
            "Las lecciones de la intervención se concentran en la Tabla 21. No se reitera aquí "
            "el recuento de incrementos ni el caso de refinamiento ya descrito en el apartado 6."
        ),
        "Se aplicó el Prompt-Centered": (
            "Se aplicó el Prompt-Centered SDLC v1.2 de forma documentada. La evidencia de prompts, "
            "evaluaciones, matriz y puesta en producción soporta el apartado 4.1. Las conclusiones "
            "y recomendaciones generales de la tesis no se enuncian en este anexo."
        ),
        "Las referencias de este anexo": (
            "A continuación se listan las fuentes citadas en el presente anexo, en el estilo de la "
            "universidad. Se incluyen las referencias del apartado 2 y las normas y autores "
            "adicionales efectivamente utilizados."
        ),
    }

    # refresh table list after deletions
    for t in list(doc.tables):
        if not is_1x1(t):
            continue
        if _table_in_block_35(t):
            continue
        text = cell0(t)
        if text.startswith(skip_starts):
            continue
        rewrite = None
        for key, val in rewrites.items():
            if text.startswith(key):
                rewrite = val
                break
        convert_1x1_to_paragraph(t, rewrite)

    # Tabla 8: quitar Cursor duplicado, fusionar Git, producción
    t8 = None
    for t in doc.tables:
        if len(t.rows) >= 6 and t.cell(0, 0).text.strip().startswith("Herramienta"):
            t8 = t
            break
    if t8 is not None:
        clear_table_keep_header(t8, [
            ["Cursor (entorno y agente)", "Generación y refinamiento de artefactos, con revisión humana", "R, D, I, T, M"],
            ["PHP 8.1+ / Composer", "API REST, PDO, JWT y dependencias del servidor", "I, T, M"],
            ["React 19 / Vite / Axios / Tailwind", "Interfaz de usuario y cliente HTTP", "I, T"],
            ["MySQL", "Almacenamiento relacional del sistema", "I, T, M"],
            ["Python 3.10+ / FastAPI / Uvicorn", "Servicio opcional de aprendizaje automático", "D, I, T"],
            ["Scikit-learn / pandas / joblib", "Entrenamiento y persistencia del modelo de riesgo", "D, I, T"],
            ["XAMPP", "Ambiente local de desarrollo y verificación", "I, T, M"],
            ["Ambiente de producción institucional", "Puesta en operación, pruebas de humo y rollback", "M"],
            ["Git / GitHub", "Versionado local y evidencia por commit", "I, T, M"],
            ["PHPUnit / Vitest / pytest", "Ejecución de pruebas; el responsable interpreta los resultados", "T"],
            ["Dompdf / jsPDF", "Exportación de fichas y reportes", "I"],
            ["Openspout", "Plantilla y carga masiva en hoja de cálculo", "I"],
            ["Bizagi", "Modelado de procesos institucionales (apartados 3.4 y 3.5)", "R"],
        ])

    # Tabla 5 fuera de alcance
    for t in doc.tables:
        if len(t.rows) >= 9 and t.cell(0, 0).text.strip() == "Campo":
            for row in t.rows:
                if "Fuera de alcance" in row.cells[0].text:
                    set_cell(
                        row.cells[1],
                        "Portal ciudadano, bienes no informáticos, SIGA/SIAF y ML obligatorio "
                        "en producción (el módulo se desactiva si el servicio predictivo no está disponible)",
                        9,
                    )
                if row.cells[0].text.strip() == "Contenido (borrador — revíselo)":
                    set_cell(row.cells[0], "Contenido", 9)

    # Tabla 13 M-001
    for t in doc.tables:
        if len(t.columns) == 4 and t.cell(0, 0).text.strip() == "Código" and "M-001" in t.cell(1, 0).text:
            set_cell(t.cell(1, 1), "Plan de puesta en producción", 8)

    # Tabla 16 DevOps: sin SHA repetidos
    for t in doc.tables:
        if len(t.columns) == 7 and "Control de versiones" in t.cell(0, 3).text:
            clear_table_keep_header(t, [
                ["F-003", "I-002", "véase Tabla 14", "Sí", "No", "Parcial", "No"],
                ["F-007", "I-007", "véase Tabla 14", "Sí", "No", "Sí (pytest)", "Parcial (ADR-002)"],
                ["F-001 a F-008", "I-002 a I-009", "véase Tabla 14", "Sí", "No", "Parcial", "No / parcial ML"],
            ])
            set_cell(t.cell(0, 0), "Función", 8)
            set_cell(t.cell(0, 1), "Prompt", 8)
            set_cell(t.cell(0, 2), "Commit", 8)
            set_cell(t.cell(0, 3), "Control de versiones", 8)
            set_cell(t.cell(0, 4), "CI", 8)
            set_cell(t.cell(0, 5), "Pruebas", 8)
            set_cell(t.cell(0, 6), "Monitoreo", 8)

    # Tabla 17: 3 filas nuevas
    for t in doc.tables:
        if len(t.columns) == 3 and t.cell(0, 0).text.strip() == "Regla" and "Revisión humana" in t.cell(1, 0).text:
            clear_table_keep_header(t, [
                ["Criterio de parada", "No se itera sin decisión D2", "Apartado 6"],
                ["Alcance del agente", "No se encadenan agentes sin revisión humana", "Apartado 8.1"],
                ["Techo de autonomía", "No se adoptan los niveles N4–N5 en producción", "Apartado 8.2"],
            ])

    # Tabla 21 lección ML
    for t in doc.tables:
        if len(t.columns) == 2 and t.cell(0, 0).text.strip() == "Hallazgo":
            for row in t.rows:
                if "Hostinger" in row.cells[0].text or "Hostinger" in row.cells[1].text or "no corre FastAPI" in row.cells[0].text:
                    set_cell(row.cells[0], "El componente de ML es opcional en producción institucional", 8)
                    set_cell(row.cells[1], "El caso no cubre un ciclo de vida de ML con servicio dedicado en operación", 8)

    # Apéndice F
    for t in doc.tables:
        if len(t.columns) == 3 and t.cell(0, 0).text.strip() == "Apéndice":
            for row in t.rows:
                if row.cells[0].text.strip() == "F":
                    set_cell(row.cells[1], "Evidencias de revisión y puesta en producción", 8)
                    set_cell(row.cells[2], "Checklist, pruebas de humo y responsable", 8)

    # Añadir referencias extra al final de la bibliografía existente
    last_bib = None
    for p in doc.paragraphs:
        if p.style and p.style.name == "Bibliography":
            last_bib = p
    extra_refs = [
        "[17]\tRojas Camayo, Prompt-Centered SDLC, versión 1.2.",
        "[18]\tISO 9001, Sistemas de gestión de la calidad. Requisitos.",
        "[19]\tISO/IEC 25010, Systems and software Quality Requirements and Evaluation (SQuaRE). Modelos de calidad.",
        "[20]\tR. Amaro, R. Pereira y M. Mira da Silva, trabajo sobre capacidades DevOps y procesos del ciclo de vida (matriz V3).",
        "[21]\tJiang et al., trabajo sobre versionado y calidad de prompts (matriz V3).",
    ]
    if last_bib is not None:
        anchor = last_bib._element
        for ref in extra_refs:
            p = insert_paragraph_after(anchor, ref, size=12, first_indent=0)
            # bibliography-like
            for run in p.runs:
                _font(run, size=11)
            anchor = p._element

    # Header
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        for run in hp.runs:
            if "Plantilla" in run.text or "guía" in run.text or "guia" in run.text:
                run.text = "Anexo 07  ·  Intervención metodológica  ·  Prompt-Centered SDLC v1.2"

    # Limpieza residual de Hostinger fuera de 3.4-3.5 (párrafos y celdas)
    residual = [
        ("Hostinger", "el ambiente de producción"),
        ("hostinger", "puesta en producción"),
        ("hPanel", "servidor institucional"),
    ]
    replace_run_text(doc, residual, skip_35=True)
    for t in doc.tables:
        if _table_in_block_35(t):
            continue
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if not run.text:
                            continue
                        txt = run.text
                        for old, new in residual:
                            if old in txt:
                                txt = txt.replace(old, new)
                        run.text = txt

    # Párrafo de nota tras Figura 1 si quedó huérfano
    p11 = find_para(doc, lambda p: p.text.strip() == "1.1 Propósito de este anexo")
    # ya insertamos purpose

    # 8.3 intro: si no quedó párrafo, el heading basta

    # Nota académica tras Tabla 14
    for p in doc.paragraphs:
        if p.text.strip().startswith("Tabla 14."):
            insert_paragraph_after(
                p._element,
                "Las evoluciones I-008 e I-009 no sustituyen el commit de origen. El apartado 7 "
                "remite a esta tabla y no reproduce los identificadores SHA.",
                size=12,
            )
            break

    # 4.3: frase breve si DoD table exists
    for p in doc.paragraphs:
        if p.text.strip() == "4.3 Definition of Done ampliada":
            insert_paragraph_after(
                p._element,
                "Los criterios de cierre de un incremento asistido por IA se presentan en la Tabla 7.",
            )
            break

    doc.save(str(DOC))
    print("OK", DOC)


if __name__ == "__main__":
    main()

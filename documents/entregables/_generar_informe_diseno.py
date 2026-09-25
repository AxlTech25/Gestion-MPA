# -*- coding: utf-8 -*-
"""Genera el informe Word de la seccion 4: Diseno y modelado del sistema."""
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

BASE = Path(__file__).resolve().parent
IMG = BASE / "_figuras_informe_diseno"
OUT = BASE / "04_Diseno_y_modelado_del_sistema.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x1A, 0x5F, 0x7A)
GRAY = RGBColor(0x4A, 0x55, 0x68)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

FONT = "Times New Roman"


def _set_run_font(run, size=12, bold=False, italic=False, color=BLACK, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, fill):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, val in kwargs.items():
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), val.get("val", "single"))
        element.set(qn("w:sz"), val.get("sz", "4"))
        element.set(qn("w:color"), val.get("color", "1F3A5F"))


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_paragraph_spacing(p, before=0, after=8, line=1.5, first_line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    _set_run_font(run, size=10, color=GRAY)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Sigemad MPA V2  ·  Intervención metodológica  ·  Diseño")
    _set_run_font(run, size=9, italic=True, color=GRAY)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = fp.add_run("Municipalidad Provincial de Acobamba")
    _set_run_font(left, size=9, color=GRAY)
    mid = fp.add_run("     —     ")
    _set_run_font(mid, size=9, color=GRAY)
    add_page_number(fp)


def add_heading_custom(doc, text, level):
    sizes = {0: 22, 1: 16, 2: 13, 3: 12}
    colors = {0: NAVY, 1: NAVY, 2: TEAL, 3: NAVY}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    before = {0: 6, 1: 18, 2: 14, 3: 10}[level]
    after = {0: 12, 1: 8, 2: 6, 3: 4}[level]
    set_paragraph_spacing(p, before=before, after=after, line=1.15, first_line=0)
    run = p.add_run(text)
    _set_run_font(run, size=sizes[level], bold=True, color=colors[level])
    if level in (0, 1):
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12" if level == 0 else "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1F3A5F" if level == 0 else "1A5F7A")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_body(doc, text, first_indent=0.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=8, line=1.5, first_line=first_indent)
    run = p.add_run(text)
    _set_run_font(run, size=12)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=4, line=1.15, first_line=0)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    bullet = p.add_run("•  ")
    _set_run_font(bullet, size=12, color=TEAL)
    if bold_lead:
        lead = p.add_run(bold_lead)
        _set_run_font(lead, size=12, bold=True)
        rest = p.add_run(text)
        _set_run_font(rest, size=12)
    else:
        run = p.add_run(text)
        _set_run_font(run, size=12)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=12, line=1.0, first_line=0)
    run = p.add_run(text)
    _set_run_font(run, size=10, italic=True, color=GRAY)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=2, after=10, line=1.15, first_line=0)
    label = p.add_run("Nota. ")
    _set_run_font(label, size=10, italic=True, bold=True, color=GRAY)
    run = p.add_run(text)
    _set_run_font(run, size=10, italic=True, color=GRAY)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = Cm(15.5)
    if col_widths is None:
        col_widths = [usable / len(headers)] * len(headers)
    for i, w in enumerate(col_widths):
        for cell in table.columns[i].cells:
            cell.width = w

    border = {"val": "single", "sz": "4", "color": "8AA0B8"}

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "1F3A5F")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        run = cell.paragraphs[0].add_run(h)
        _set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        set_cell_border(cell, top=border, bottom=border, left=border, right=border)

    for r_idx, row in enumerate(rows):
        prevent_row_split(table.rows[r_idx + 1])
        fill = "F4F7FB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            shade_cell(cell, fill)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            run = cell.paragraphs[0].add_run(str(value))
            _set_run_font(run, size=10)
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    return table


def add_picture_centered(doc, path, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=6, after=2, line=1.0, first_line=0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def _box(ax, x, y, w, h, text, facecolor, edge="#1F3A5F", fontsize=9, sub=None):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.02,rounding_size=0.08",
        linewidth=1.4, edgecolor=edge, facecolor=facecolor,
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2 + (0.08 if sub else 0), text,
            ha="center", va="center", fontsize=fontsize, fontweight="bold",
            color="#1A1A1A", wrap=True)
    if sub:
        ax.text(x + w / 2, y + h / 2 - 0.22, sub,
                ha="center", va="center", fontsize=7.2, color="#4A5568")


def _arrow(ax, x1, y1, x2, y2, text=None):
    ax.annotate(
        "", xy=(x2, y2), xytext=(x1, y1),
        arrowprops=dict(arrowstyle="-|>", color="#1F3A5F", lw=1.3),
    )
    if text:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        ax.text(mx + 0.12, my, text, fontsize=7, color="#1A5F7A", fontstyle="italic")


def draw_arquitectura():
    IMG.mkdir(exist_ok=True)
    path = IMG / "figura1_arquitectura_general.png"
    fig, ax = plt.subplots(figsize=(10.2, 6.4), dpi=180)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6.4)
    ax.axis("off")
    fig.patch.set_facecolor("white")

    _box(ax, 2.4, 5.35, 5.2, 0.85, "Usuario  /  Navegador web", "#E8EEF6")
    _arrow(ax, 5.0, 5.35, 5.0, 4.85)
    _box(ax, 1.6, 3.85, 6.8, 1.0, "Capa de presentación — SPA React (Vite)", "#D6E4F0",
         sub="Axios  ·  JWT  ·  React Router  ·  features por dominio")
    _arrow(ax, 5.0, 3.85, 5.0, 3.35, "  HTTP + JSON + JWT")
    _box(ax, 1.6, 2.25, 6.8, 1.1, "Capa de aplicación — API REST PHP 8", "#B8CDE0",
         sub="Front controller  ·  MVC  ·  middleware JWT / RBAC  ·  /api/v2")
    _arrow(ax, 3.4, 2.25, 2.5, 1.55)
    _arrow(ax, 6.6, 2.25, 7.5, 1.55)
    _box(ax, 0.45, 0.25, 4.1, 1.3, "Capa de datos — MySQL", "#C5DDD4",
         sub="BD gestion_equipos_mpa_v2   tablas v2_*")
    _box(ax, 5.45, 0.25, 4.1, 1.3, "Microservicio ML (opcional)", "#F3E4C8",
         sub="FastAPI :8000  ·  Scikit-learn  ·  no expuesto al navegador")
    ax.annotate(
        "", xy=(6.4, 0.9), xytext=(3.6, 0.9),
        arrowprops=dict(arrowstyle="<->", color="#8A6D3B", lw=1.1),
    )
    ax.text(5.0, 1.12, "proxy PHP", fontsize=7, ha="center", color="#8A6D3B")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def draw_mvc():
    path = IMG / "figura2_flujo_mvc.png"
    fig, ax = plt.subplots(figsize=(10.2, 3.8), dpi=180)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3.6)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    items = [
        (0.2, "Petición\nHTTP", "#E8EEF6"),
        (2.15, "index.php\nCORS + JWT", "#D6E4F0"),
        (4.1, "routes/\nrecurso", "#C5D8EA"),
        (6.05, "Controller", "#B8CDE0"),
        (8.0, "Model\nPDO → MySQL", "#C5DDD4"),
    ]
    for x, label, color in items:
        _box(ax, x, 1.35, 1.75, 1.15, label, color, fontsize=8.2)
    for x in (1.95, 3.9, 5.85, 7.8):
        _arrow(ax, x, 1.92, x + 0.2, 1.92)
    ax.text(5.0, 0.55, "Respuesta JSON  { success, data, message }",
            ha="center", fontsize=8, color="#1A5F7A", fontstyle="italic")
    ax.annotate("", xy=(0.95, 1.35), xytext=(8.85, 1.15),
                arrowprops=dict(arrowstyle="-|>", color="#1A5F7A",
                                connectionstyle="arc3,rad=-0.28", lw=1.1))
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def draw_er_conceptual():
    path = IMG / "figura3_modelo_conceptual.png"
    fig, ax = plt.subplots(figsize=(10.4, 7.2), dpi=180)
    ax.set_xlim(0, 10.4)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")

    def ebox(x, y, w, h, title, color="#E8EEF6"):
        _box(ax, x, y, w, h, title, color, fontsize=8)

    ebox(0.3, 5.9, 2.2, 0.75, "ÁREA", "#C5DDD4")
    ebox(3.2, 5.9, 2.2, 0.75, "USUARIO", "#C5DDD4")
    ebox(4.0, 3.85, 2.5, 0.9, "EQUIPO", "#B8CDE0")
    ebox(7.5, 3.95, 2.5, 0.75, "FICHA TÉCNICA", "#D6E4F0")
    ebox(0.25, 2.15, 2.5, 0.75, "CATEGORÍA FALLA", "#C5DDD4")
    ebox(3.7, 2.05, 2.9, 0.85, "FICHA MANTENIMIENTO", "#F3E4C8")
    ebox(7.3, 2.15, 2.7, 0.75, "PREDICCIÓN ML", "#F3E4C8")
    ebox(0.25, 0.35, 2.3, 0.7, "HISTORIAL ASIG.", "#E8EEF6")
    ebox(2.75, 0.35, 2.3, 0.7, "CRONOGRAMA", "#E8EEF6")
    ebox(5.25, 0.35, 2.3, 0.7, "HOJA DE BAJA", "#E8EEF6")
    ebox(7.75, 0.35, 2.3, 0.7, "MÉTRICA", "#F3E4C8")

    def link(x1, y1, x2, y2, label, dx=0.0, dy=0.08):
        ax.plot([x1, x2], [y1, y2], color="#1F3A5F", lw=1.05)
        ax.text((x1 + x2) / 2 + dx, (y1 + y2) / 2 + dy, label,
                fontsize=6.4, color="#1A5F7A", ha="center")

    link(2.5, 6.28, 3.2, 6.28, "1 : N")
    link(4.3, 5.9, 5.0, 4.75, "1 : N  responsable", dx=0.55)
    link(1.4, 5.9, 4.4, 4.75, "1 : N  asigna", dx=-0.7)
    link(6.5, 4.3, 7.5, 4.3, "1 : 1")
    link(5.25, 3.85, 5.15, 2.9, "1 : N")
    link(2.75, 2.52, 3.7, 2.47, "1 : N")
    link(6.6, 2.47, 7.3, 2.52, "1 : N")
    link(4.5, 3.85, 1.4, 1.05, "", dy=0)
    link(5.25, 3.85, 3.9, 1.05, "", dy=0)
    link(5.25, 3.85, 6.4, 1.05, "", dy=0)
    link(5.25, 3.85, 8.9, 1.05, "1 : N", dx=0.85, dy=0.15)
    link(5.15, 2.05, 8.9, 1.05, "1 : N", dx=0.35, dy=-0.12)
    link(4.3, 5.9, 5.15, 2.9, "técnico  1 : N", dx=1.15, dy=0)

    ax.text(5.2, 6.95, "Entidad central: EQUIPO   ·   Catálogos en verde   ·   Históricos / ML en ámbar",
            ha="center", fontsize=7.5, color="#4A5568", fontstyle="italic")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def build():
    fig1 = draw_arquitectura()
    fig2 = draw_mvc()
    fig3 = draw_er_conceptual()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    configure_section(doc.sections[0])

    # --- Portada ---
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=4, line=1.15, first_line=0)
    r = p.add_run("INTERVENCIÓN METODOLÓGICA")
    _set_run_font(r, size=13, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=18, line=1.15, first_line=0)
    r = p.add_run("Prompt-Centered SDLC v1.2  ·  Caso de estudio")
    _set_run_font(r, size=11, italic=True, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=8, line=1.15, first_line=0)
    r = p.add_run("4.  Diseño y modelado del sistema")
    _set_run_font(r, size=26, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=22, line=1.15, first_line=0)
    r = p.add_run("(software)")
    _set_run_font(r, size=18, italic=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=6, line=1.15, first_line=0)
    r = p.add_run("Sigemad MPA V2")
    _set_run_font(r, size=16, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=4, line=1.15, first_line=0)
    r = p.add_run("Sistema de gestión de equipos de cómputo\ny mantenimiento predictivo")
    _set_run_font(r, size=12, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=16, after=4, line=1.15, first_line=0)
    r = p.add_run("Municipalidad Provincial de Acobamba")
    _set_run_font(r, size=12, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=4, line=1.15, first_line=0)
    r = p.add_run("Unidad de informática / patrimonio de TI")
    _set_run_font(r, size=11, color=GRAY)

    meta = [
        ("Producto", "SIGEMAD-MPA-V2"),
        ("Versión del software", "0.9.0"),
        ("Documento de diseño", "ADR-001  ·  architecture.md"),
        ("Base de datos", "gestion_equipos_mpa_v2  (MySQL, prefijo v2_)"),
        ("Fecha del informe", "septiembre de 2026"),
    ]
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=28, after=8, first_line=0)
    add_table(doc, ["Campo", "Contenido"], meta, [Cm(5.2), Cm(10.3)])

    doc.add_page_break()

    # --- Indice ---
    add_heading_custom(doc, "Contenido", 1)
    toc = [
        "4.     Diseño y modelado del sistema (software)",
        "4.1.   Arquitectura del software (por capas, microservicios, APIs)",
        "4.1.1. Arquitectura del proyecto en general",
        "4.1.2. Arquitectura interna (equivalente al MVC del framework)",
        "4.2.   Diseño de base de datos (modelo entidad-relación y normalización)",
        "4.2.1. Modelo conceptual",
        "4.2.2. Normalización en fase 1",
        "4.2.3. Modelo lógico",
        "4.2.4. Normalización en fase 2",
        "4.2.5. Modelo físico (entidad-relación implementado)",
        "4.3.   Síntesis para la defensa",
        "Referencias de diseño del repositorio",
    ]
    for item in toc:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=3, line=1.3, first_line=0)
        r = p.add_run(item)
        _set_run_font(r, size=12)

    doc.add_page_break()

    # --- 4 ---
    add_heading_custom(doc, "4.  Diseño y modelado del sistema (software)", 1)
    add_body(
        doc,
        "En esta etapa se definió cómo se construye el producto: la forma de las capas, "
        "la comunicación entre componentes y el modelo de datos que sostiene el inventario "
        "patrimonial, las fichas técnicas, el mantenimiento y la predicción de riesgo. "
        "El diseño no parte de un framework único tipo Laravel, sino de una decisión "
        "documentada en el ADR-001: una aplicación de página única (SPA) en React, una "
        "API REST en PHP y una base de datos MySQL, complementadas por un microservicio "
        "de aprendizaje automático opcional en FastAPI."
    )
    add_body(
        doc,
        "El criterio de diseño fue doble. Primero, el sistema debe operar en un entorno "
        "local con XAMPP y, en producción, en un servidor Apache + PHP + MySQL. Segundo, "
        "los datos debían nacer estructurados y numéricos —y no como texto libre—, para que "
        "el componente de machine learning no dependiera de descripciones inconsistentes. "
        "Esa restricción atraviesa tanto la arquitectura de software como el modelado de la "
        "base de datos."
    )

    # --- 4.1 ---
    add_heading_custom(doc, "4.1.  Arquitectura del software (por capas, microservicios, APIs)", 1)
    add_body(
        doc,
        "La arquitectura se describe en dos niveles, tal como se exige en esta intervención: "
        "la del producto completo (qué sistemas conversan y qué servicios externos se utilizan) "
        "y la interna de cada lado (el equivalente al MVC de tres capas de un framework, "
        "aunque en este caso no se emplea Laravel)."
    )

    add_heading_custom(doc, "4.1.1.  Arquitectura del proyecto en general", 2)
    add_body(
        doc,
        "Sigemad MPA V2 adopta una arquitectura cliente-servidor de tres capas, con un "
        "monolito modular en PHP y un sidecar de machine learning que puede apagarse sin "
        "tumbar el resto del sistema. La migración desde la versión 1 se realizó con el "
        "patrón Strangler Fig: las tablas nuevas llevan el prefijo v2_ y conviven con el "
        "esquema anterior hasta reemplazarlo por completo."
    )

    add_picture_centered(doc, fig1, 15.4)
    add_caption(doc, "Figura 1. Arquitectura general de Sigemad MPA V2: tres capas más un microservicio ML opcional.")

    add_body(
        doc,
        "El navegador nunca llama a Python. PHP actúa como proxy autenticado (clase MlService): "
        "recibe el token JWT, valida el rol y, si el servicio está configurado, reenvía la "
        "petición a FastAPI. Si la variable ml_service_url está vacía —cuando FastAPI no "
        "está desplegado— o si FastAPI no responde, el inventario, las fichas y el "
        "mantenimiento siguen disponibles; solo se omiten los badges de riesgo y las "
        "sugerencias de categoría de falla. Esa degradación controlada fue una restricción "
        "explícita del diseño (prompt D-002 y ADR-001).",
        first_indent=0.75,
    )

    add_body(doc, "Servicios y dependencias que el proyecto utiliza hacia afuera:", first_indent=0)
    add_table(
        doc,
        ["Componente", "Tipo", "Función", "Obligatorio"],
        [
            ["Apache + PHP 8.1+", "Servidor de aplicación\n(XAMPP / producción)", "Expone la API y, en producción, el frontend compilado", "Sí"],
            ["MySQL", "Gestor de base de datos", "Persistencia de inventario, usuarios, fichas y predicciones", "Sí"],
            ["FastAPI (Python 3.10+)", "Microservicio interno", "Predicción de riesgo y sugerencia de categoría de falla", "No"],
            ["Cliente HTTP Axios", "Consumo de la API propia", "El navegador solo habla con PHP", "Sí"],
        ],
        [Cm(4.0), Cm(3.6), Cm(5.3), Cm(2.6)],
    )
    add_caption(doc, "Tabla 1. Servicios externos e internos del producto.")
    add_note(
        doc,
        "No existen APIs de terceros de negocio (pasarelas de pago, OAuth social o nubes de ML). "
        "Las librerías firebase/php-jwt, Dompdf, OpenSpout y jsPDF se ejecutan dentro del propio "
        "código; no constituyen servicios remotos."
    )

    add_body(doc, "Reglas de integración adoptadas:", first_indent=0)
    add_bullet(
        doc,
        " El navegador no tiene ruta directa al puerto 8000. Toda inferencia pasa por /api/v2/ml/*.",
        bold_lead="Proxy autenticado.",
    )
    add_bullet(
        doc,
        " Toda ruta de /api/v2/* excepto /auth exige JWT. La interfaz (PrivateRoute y RoleRoute) "
        "oculta pantallas; la autorización efectiva está en el servidor (RBAC: Administrador, "
        "Técnico, Practicante).",
        bold_lead="Defensa en profundidad.",
    )
    add_bullet(
        doc,
        " Si FastAPI no está desplegado, el ML se deshabilita sin recompilar el núcleo operativo.",
        bold_lead="Degradación elegante.",
    )

    add_body(
        doc,
        "El estilo de la API es REST. Las respuestas siguen un contrato uniforme "
        "{ success, data, message }. Los recursos principales son /auth, /equipos, "
        "/mantenimientos, /fichas-tecnicas, /dashboard, /reportes, /areas, /usuarios y /ml.",
    )

    add_heading_custom(doc, "4.1.2.  Arquitectura interna (equivalente al MVC del framework)", 2)
    add_body(
        doc,
        "En un proyecto Laravel se hablaría de la arquitectura en tres capas MVC que el "
        "framework impone por convención (rutas, controladores, modelos Eloquent y vistas Blade). "
        "Sigemad MPA no adopta Laravel: el backend es PHP plano orientado a objetos con PDO y el "
        "frontend es React. Aun así, cada lado replica una arquitectura de capas comparable, "
        "elegida a propósito para un hosting PHP compartido y un equipo pequeño (alternativa "
        "descartada en el ADR-001: monolito Laravel + Vue, usado en SGMI, no en este producto)."
    )

    add_heading_custom(doc, "A.  Backend PHP — MVC adaptado a API REST", 3)
    add_table(
        doc,
        ["Capa MVC", "Ubicación", "Responsabilidad"],
        [
            ["Vista / contrato HTTP", "routes/  +  JSON (y PDF)", "No hay Blade ni Twig: la vista es el JSON de la API y los reportes PDF"],
            ["Controlador", "controllers/", "Interpreta la petición, aplica reglas de rol y orquesta el caso de uso"],
            ["Modelo", "models/", "Acceso a MySQL mediante PDO orientado a objetos"],
            ["Transversal", "middleware/, config/, services/", "JWT, CORS, conexión a BD y cliente del microservicio ML"],
        ],
        [Cm(4.2), Cm(4.4), Cm(6.9)],
    )
    add_caption(doc, "Tabla 2. Correspondencia entre el MVC clásico y la API V2.")

    add_body(
        doc,
        "El punto de entrada es el front controller backend/api/v2/index.php: aplica CORS, "
        "exige autenticación salvo en /auth, despacha el recurso y carga la ruta correspondiente. "
        "El flujo es el mismo que en Laravel (ruta → controlador → modelo), sin el contenedor "
        "de inversiones de control del framework.",
        first_indent=0,
    )

    add_picture_centered(doc, fig2, 15.4)
    add_caption(doc, "Figura 2. Flujo de una petición en la API V2 (MVC adaptado a REST).")

    add_heading_custom(doc, "B.  Frontend React — arquitectura por features (dominios)", 3)
    add_body(
        doc,
        "React no es un MVC de servidor. El código se organizó por módulos de negocio y no "
        "por tipo de archivo, de modo que cada dominio agrupa sus pantallas y su cliente HTTP:",
        first_indent=0,
    )
    add_bullet(doc, " login y token JWT.", bold_lead="features/auth —")
    add_bullet(doc, " equipos, ficha técnica y carga Excel.", bold_lead="features/inventario —")
    add_bullet(doc, " intervenciones e historial.", bold_lead="features/mantenimiento —")
    add_bullet(doc, " indicadores y consulta por etiquetas.", bold_lead="features/dashboard —")
    add_bullet(doc, " áreas y personal (solo Administrador).", bold_lead="features/configuracion —")
    add_bullet(doc, " badges de riesgo y consumo del proxy PHP.", bold_lead="features/ml —")
    add_body(
        doc,
        "Encima hay capas compartidas: components/ (barra de navegación), context/ (sesión) "
        "y lib/ (cliente Axios). El enrutador App.jsx cumple el rol de controlador de "
        "navegación. Las rutas privadas usan PrivateRoute (sesión) y RoleRoute (rol).",
    )

    add_heading_custom(doc, "C.  Microservicio ML — capas propias", 3)
    add_body(
        doc,
        "FastAPI separa routers (HTTP), services (entrenamiento e inferencia) y schemas. "
        "Los modelos serializados (.joblib) viven fuera de la API PHP. Se trata de un "
        "microservicio acotado a la predicción, no de una malla de muchos servicios. "
        "En términos de estilo arquitectónico, el producto es un monolito modular con un "
        "sidecar opcional, no una arquitectura de microservicios completa.",
        first_indent=0,
    )

    # --- 4.2 ---
    add_heading_custom(doc, "4.2.  Diseño de base de datos (modelo entidad-relación y normalización)", 1)
    add_body(
        doc,
        "La base gestion_equipos_mpa_v2 se diseñó siguiendo las tres fases del modelado "
        "enseñadas en el curso de Bases de Datos —conceptual, lógico y físico— y dos pasadas "
        "de normalización. El criterio adicional de este producto es no usar texto libre como "
        "fuente de verdad en áreas ni en categorías de falla, para poder entrenar el modelo "
        "de riesgo (prompt D-001)."
    )

    add_heading_custom(doc, "4.2.1.  Modelo conceptual", 2)
    add_body(
        doc,
        "El modelo conceptual responde a qué existe en el dominio y cómo se relaciona, "
        "sin declarar aún tipos de dato. Las entidades identificadas son las siguientes.",
    )
    add_table(
        doc,
        ["Entidad", "Qué representa"],
        [
            ["Área", "Unidad organizacional (oficina o gerencia) a la que se asignan equipos y personal"],
            ["Usuario", "Persona del sistema con rol Administrador, Técnico o Practicante"],
            ["Equipo", "Bien patrimonial de cómputo (entidad central del modelo)"],
            ["Ficha técnica", "Especificación uno a uno del equipo (hardware, software y evaluación)"],
            ["Categoría de falla", "Catálogo de tipos de problema, con severidad"],
            ["Ficha de mantenimiento", "Intervención histórica sobre un equipo"],
            ["Historial de asignaciones", "Traslados de área o de responsable en el tiempo"],
            ["Cronograma de mantenimiento", "Programación preventiva"],
            ["Hoja de baja", "Revisión técnica para retiro, excedencia o chatarra"],
            ["Predicción ML", "Score de riesgo generado por una versión del modelo"],
            ["Métrica de equipo", "Lectura de telemetría en un instante"],
        ],
        [Cm(5.2), Cm(10.3)],
    )
    add_caption(doc, "Tabla 3. Entidades del modelo conceptual.")

    add_body(doc, "Cardinalidades entre entidades:", first_indent=0)
    add_table(
        doc,
        ["Relación", "Cardinalidad", "Significado"],
        [
            ["Área — Usuario", "1 : N", "Un área agrupa a varios usuarios"],
            ["Área — Equipo", "1 : N", "Un área tiene varios equipos vigentes"],
            ["Usuario — Equipo", "1 : N", "Un usuario puede ser responsable de varios equipos"],
            ["Equipo — Ficha técnica", "1 : 1", "Cada equipo tiene una única ficha de especificación"],
            ["Equipo — Ficha de mantenimiento", "1 : N", "Un equipo acumula muchas intervenciones"],
            ["Categoría de falla — Ficha de mant.", "1 : N", "Una categoría clasifica muchas fichas"],
            ["Usuario — Ficha de mantenimiento", "1 : N", "Un técnico registra varias intervenciones"],
            ["Equipo — Historial / cronograma / baja", "1 : N", "Hechos históricos o de planificación"],
            ["Equipo — Predicción ML / métrica", "1 : N", "Series de riesgo y de telemetría"],
            ["Ficha de mant. — Métrica", "1 : N", "Lecturas asociadas a una intervención"],
        ],
        [Cm(5.6), Cm(2.6), Cm(7.3)],
    )
    add_caption(doc, "Tabla 4. Cardinalidades del modelo conceptual.")

    add_picture_centered(doc, fig3, 15.5)
    add_caption(doc, "Figura 3. Modelo conceptual: el equipo como entidad central, catálogos y hechos históricos.")

    add_body(
        doc,
        "En síntesis conceptual, el Equipo es el núcleo. Área y Categoría de falla son "
        "catálogos. La ficha técnica es una extensión uno a uno. El resto de entidades "
        "registran historia operativa o resultados del modelo predictivo.",
    )

    add_heading_custom(doc, "4.2.2.  Normalización en fase 1", 2)
    add_body(
        doc,
        "La primera pasada de normalización separa lo independiente, lo dependiente y lo que "
        "se repetía como texto en varias tablas (catálogos que en la versión 1 vivían incrustados "
        "en el registro del equipo o de la ficha)."
    )

    add_heading_custom(doc, "Tablas independientes (maestras o catálogos)", 3)
    add_body(
        doc,
        "No necesitan de otra entidad para existir. Área: en la V1 el nombre del área se "
        "escribía en cada equipo; ahora es la tabla v2_areas. Categoría de falla: en la V1 el "
        "diagnóstico era prosa libre; ahora es el catálogo v2_categorias_falla, con severidad "
        "(Baja, Media, Alta, Crítica). Este desglose es imprescindible para el aprendizaje "
        "automático, porque un clasificador no puede aprender de sinónimos escritos a mano.",
        first_indent=0,
    )

    add_heading_custom(doc, "Tablas dependientes (fuertes)", 3)
    add_body(
        doc,
        "Existen porque hay un catálogo o un dueño. Usuario depende de Área (area_id). "
        "Equipo depende de Área de forma obligatoria y, de forma opcional, de Usuario "
        "(responsable actual).",
        first_indent=0,
    )

    add_heading_custom(doc, "Tablas dependientes débiles o de detalle", 3)
    add_body(
        doc,
        "No tienen sentido de negocio sin el equipo (o sin la intervención): ficha técnica, "
        "ficha de mantenimiento, historial de asignaciones, cronograma, hoja de baja, "
        "predicción ML y métrica de equipo.",
        first_indent=0,
    )

    add_heading_custom(doc, "Atributos repetitivos o «polimórficos» que se desglosaron", 3)
    add_body(
        doc,
        "En el sentido del curso, se trata de valores que se copiaban en muchas filas y se "
        "extrajeron a una tabla propia. No hay tablas polimórficas clásicas (un campo "
        "entidad_tipo más entidad_id). Lo que se evitó fue la repetición de catálogos.",
        first_indent=0,
    )
    add_table(
        doc,
        ["Antes (V1 / texto repetido)", "Después (V2, normalizado)"],
        [
            ["Nombre de área escrito en cada equipo", "v2_areas + equipo.area_id"],
            ["Tipo de falla en prosa en cada ficha", "v2_categorias_falla + categoria_falla_id"],
            ["Rol y estados como cadenas sueltas", "Dominios ENUM controlados (rol, estado operativo, tipo de equipo)"],
        ],
        [Cm(7.6), Cm(7.9)],
    )
    add_caption(doc, "Tabla 5. Desglose de atributos repetitivos en la normalización de fase 1.")
    add_note(
        doc,
        "Quedó una denormalización consciente: responsable_nombre en el equipo, copia del "
        "nombre del usuario, para no romper listados si se elimina la cuenta (ON DELETE SET NULL "
        "sobre responsable_id). Con este paso se alcanza la primera forma normal: valores "
        "atómicos y sin grupos repetitivos de áreas o fallas."
    )

    add_heading_custom(doc, "4.2.3.  Modelo lógico", 2)
    add_body(
        doc,
        "En el modelo lógico aparecen claves, dominios y tipos. Los nombres técnicos llevan "
        "el prefijo v2_. A continuación se resume la estructura de las tablas núcleo."
    )

    add_heading_custom(doc, "v2_areas", 3)
    add_table(
        doc,
        ["Atributo", "Tipo lógico", "Restricción"],
        [
            ["id", "Entero autonumérico", "Clave primaria"],
            ["nombre", "Cadena (100)", "Único, obligatorio"],
            ["jefe_encargado", "Cadena (100)", "Opcional"],
            ["descripcion", "Texto", "Opcional"],
        ],
        [Cm(4.5), Cm(5.0), Cm(6.0)],
    )
    add_caption(doc, "Tabla 6. Estructura lógica de v2_areas.")

    add_heading_custom(doc, "v2_usuarios", 3)
    add_table(
        doc,
        ["Atributo", "Tipo lógico", "Restricción"],
        [
            ["id", "Entero autonumérico", "Clave primaria"],
            ["nombre_completo", "Cadena (100)", "Obligatorio"],
            ["usuario", "Cadena (50)", "Único, obligatorio"],
            ["password_hash", "Cadena (255)", "Hash bcrypt, no texto plano"],
            ["rol", "Dominio {Administrador, Tecnico, Practicante}", "Obligatorio"],
            ["area_id", "Entero", "FK a v2_areas, opcional"],
        ],
        [Cm(4.5), Cm(6.4), Cm(4.6)],
    )
    add_caption(doc, "Tabla 7. Estructura lógica de v2_usuarios.")

    add_heading_custom(doc, "v2_equipos", 3)
    add_table(
        doc,
        ["Atributo", "Tipo lógico", "Nota de diseño"],
        [
            ["id", "Entero", "Clave primaria"],
            ["codigo_patrimonial", "Cadena (50)", "Único; regla de negocio: 12 dígitos"],
            ["tipo_equipo", "Dominio {Laptop, CPU, Impresora, Monitor, Otro}", "Clasificación operativa y de ML"],
            ["ram_gb, almacenamiento_gb", "Entero", "Numérico para ML (no «8 GB»)"],
            ["tipo_disco", "Dominio {HDD, SSD, NVMe}", "Feature del modelo de riesgo"],
            ["horas_uso, errores_smart, contador_paginas", "Entero", "Telemetría (instantánea vigente)"],
            ["salud_bateria, ultima_temp_cpu, ultima_temp_disco", "Decimal (5,2)", "Lecturas de la última evaluación"],
            ["fecha_adquisicion", "Fecha", "Permite calcular antigüedad en meses"],
            ["costo_estimado", "Decimal (10,2)", "Dato patrimonial"],
            ["area_id", "Entero", "FK obligatoria a v2_areas"],
            ["responsable_id", "Entero", "FK opcional a v2_usuarios"],
            ["estado_conservacion / estado_operativo", "Dominios ENUM", "Estados controlados, no texto libre"],
            ["fecha_ultimo_mantenimiento", "Fecha", "Apoyo a alertas y al dataset ML"],
        ],
        [Cm(5.8), Cm(4.6), Cm(5.1)],
    )
    add_caption(doc, "Tabla 8. Estructura lógica de v2_equipos (atributos principales).")

    add_heading_custom(doc, "Resto de tablas del modelo lógico", 3)
    add_table(
        doc,
        ["Tabla", "Claves y tipos esenciales"],
        [
            ["v2_fichas_tecnicas", "id PK; equipo_id UNIQUE FK (relación 1:1); numero_ficha único; procesador, SO (ENUM), MAC, IP, observaciones, imágenes"],
            ["v2_categorias_falla", "id PK; nombre; severidad {Baja, Media, Alta, Critica}"],
            ["v2_fichas_mantenimiento", "id PK; nro_orden; equipo_id FK; tecnico_id FK; categoria_falla_id FK; tipo {Preventivo, Correctivo, Predictivo, Evaluacion}; telemetría de la visita; costo Decimal(10,2); estado_post_mantenimiento"],
            ["v2_historial_asignaciones", "id PK; equipo_id FK; área y responsable de origen/destino; tipo_movimiento ENUM; fecha_movimiento"],
            ["v2_cronograma_mantenimiento", "id PK; equipo_id FK; frecuencia_dias; proxima_fecha; prioridad; estado de programación"],
            ["v2_hojas_baja", "id PK; numero_ficha único; equipo_id FK; causales booleanas; creado_por y validado_por FK"],
            ["v2_predicciones_ml", "id PK; equipo_id FK; modelo_version; score_riesgo Decimal(5,2); nivel_riesgo ENUM; factores_json; categoria_sugerida_id FK"],
            ["v2_metricas_equipo", "id PK; equipo_id FK; mantenimiento_id FK opcional; horas, SMART, temperaturas, polvo; registrado_en"],
        ],
        [Cm(5.4), Cm(10.1)],
    )
    add_caption(doc, "Tabla 9. Resto del modelo lógico.")
    add_body(
        doc,
        "Reglas lógicas de integridad: un equipo tiene un área vigente y una ficha técnica; "
        "puede acumular muchas intervenciones y muchas predicciones. Las predicciones no "
        "alteran el inventario: solo anotan un score asociado a una versión de modelo.",
    )

    add_heading_custom(doc, "4.2.4.  Normalización en fase 2", 2)
    add_body(
        doc,
        "La segunda pasada resuelve las relaciones muchos a muchos. En este dominio no "
        "aparecen pivotes vacíos (una tabla equipo_categoria con solo dos identificadores). "
        "La relación N:N tiene atributos propios —fecha, costo, síntoma, score—, de modo que "
        "se materializa como tabla intermedia con significado de negocio."
    )
    add_table(
        doc,
        ["Relación N:N conceptual", "Tabla intermedia", "Atributos propios (por eso no es un pivote vacío)"],
        [
            ["Equipo N:N Categoría de falla", "v2_fichas_mantenimiento", "Fecha, tipo, síntoma, costo, estado posterior"],
            ["Equipo N:N Usuario (como técnico)", "v2_fichas_mantenimiento", "Quién intervino y qué hizo"],
            ["Equipo N:N Área (en el tiempo)", "v2_historial_asignaciones", "Origen, destino, motivo y fecha"],
            ["Equipo N:N lectura de sensores", "v2_metricas_equipo", "Instantánea de telemetría"],
            ["Equipo N:N versión de modelo ML", "v2_predicciones_ml", "Score, nivel y JSON de factores"],
        ],
        [Cm(4.8), Cm(4.4), Cm(6.3)],
    )
    add_caption(doc, "Tabla 10. Tablas intermedias de la normalización de fase 2.")
    add_body(
        doc,
        "La asignación actual Equipo–Área y Equipo–Responsable se dejó en N:1 (area_id, "
        "responsable_id) para que las consultas del inventario sean directas. El historial "
        "guarda el muchos a muchos temporal. Con ello el esquema queda, en lo esencial, en "
        "tercera forma normal: cada hecho no clave depende de la clave de su tabla y no de "
        "otra entidad."
    )
    add_body(doc, "Excepciones deliberadas (desnormalización controlada):", first_indent=0)
    add_bullet(
        doc,
        " instantánea de telemetría en v2_equipos (última temperatura, horas de uso) además "
        "de la serie en v2_metricas_equipo, para acelerar el dashboard y las features del modelo.",
        bold_lead="Snapshot vigente:",
    )
    add_bullet(
        doc,
        " v2_ml_equipos_features no es una tabla base; agrega correctivos a doce meses, "
        "antigüedad y severidad máxima. Es el dataset lógico del modelo, no una violación "
        "de la 3FN en el esquema operativo.",
        bold_lead="Vista de extracción:",
    )

    add_heading_custom(doc, "4.2.5.  Modelo físico (entidad-relación implementado)", 2)
    add_body(
        doc,
        "El modelo físico es el esquema ya creado en el gestor MySQL: motor InnoDB, "
        "juego de caracteres utf8mb4, claves primarias autonuméricas, claves foráneas e "
        "índices. Es el diagrama que phpMyAdmin muestra como diseñador o como vista de "
        "relaciones: el modelo implementado, no el bosquejo conceptual. Se despliega en "
        "XAMPP (desarrollo) y en phpMyAdmin del servidor de producción."
    )
    add_body(doc, "Scripts de implementación (fuente de verdad del modelo físico):", first_indent=0)
    add_table(
        doc,
        ["Script", "Qué materializa"],
        [
            ["backend/sql/v2_estructura.sql", "Núcleo: áreas, usuarios, equipos, fichas, catálogo de fallas, historial, cronograma y bajas"],
            ["backend/sql/v2_extension_fase7.sql", "Telemetría en equipos, campos estructurados de la intervención y numeración de fichas"],
            ["backend/sql/v2_ml_predicciones.sql", "Tabla de predicciones del modelo"],
            ["backend/sql/v2_metricas_equipo.sql", "Serie temporal de lecturas por equipo"],
            ["backend/sql/v2_ml_dataset_view.sql", "Vista v2_ml_equipos_features para el dataset A"],
        ],
        [Cm(6.4), Cm(9.1)],
    )
    add_caption(doc, "Tabla 11. Scripts que constituyen el modelo físico.")

    add_body(doc, "Integridad referencial implementada en el gestor:", first_indent=0)
    add_bullet(
        doc,
        " al borrar un equipo se eliminan su ficha técnica, mantenimientos, predicciones y métricas.",
        bold_lead="ON DELETE CASCADE:",
    )
    add_bullet(
        doc,
        " al borrar un usuario, las fichas conservan el hecho histórico (tecnico_id queda nulo); "
        "no se borra la intervención.",
        bold_lead="ON DELETE SET NULL:",
    )
    add_bullet(
        doc,
        " equipo_id, fecha_movimiento, proxima_fecha, nivel_riesgo, horas_uso y errores_smart, "
        "orientados a reportes y a consultas del modelo.",
        bold_lead="Índices:",
    )
    add_bullet(
        doc,
        " codigo_patrimonial, usuario, nombre de área y numero_ficha.",
        bold_lead="Unicidad:",
    )
    add_body(
        doc,
        "En el gestor, el diagrama entidad-relación físico es exactamente el conjunto de "
        "tablas v2_* unidas por esas claves foráneas. Esa es la evidencia de que el modelo "
        "conceptual y el lógico no se quedaron en el papel: están creados, con motor "
        "transaccional, y son los que usa la API V2 en tiempo de ejecución."
    )

    # --- 4.3 ---
    add_heading_custom(doc, "4.3.  Síntesis para la defensa", 1)
    add_body(
        doc,
        "La arquitectura de Sigemad MPA V2 es de tres capas con API REST y un microservicio "
        "de machine learning opcional. El backend replica el MVC sin Laravel; el frontend se "
        "organiza por dominios de negocio. La base de datos se diseñó en modelo conceptual, "
        "lógico y físico, con normalización de catálogos en la fase 1 y tablas intermedias "
        "con atributos para las relaciones muchos a muchos históricas en la fase 2. El dato "
        "sirve a la vez a la operación municipal y al modelo predictivo, y el sistema sigue "
        "operando si el servicio de inteligencia no está disponible."
    )

    add_heading_custom(doc, "Referencias de diseño del repositorio", 1)
    refs = [
        "ADR-001. Stack y arquitectura V2. documents/02_diseno/adr/ADR-001-stack-arquitectura.md",
        "Arquitectura V2. documents/02_diseno/architecture.md",
        "Prompt D-001. Modelo de datos V2. prompts/02_diseno/D-001_modelo_datos_v1.md",
        "Prompt D-002. Decisión de arquitectura. prompts/02_diseno/D-002_decision_arquitectura_v1.md",
        "Ficha de proyecto SIGEMAD-MPA-V2. documents/01_requisitos/ficha-proyecto.md",
        "Esquema físico. backend/sql/v2_estructura.sql y scripts complementarios v2_*.sql",
    ]
    for i, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_spacing(p, after=4, line=1.15, first_line=0)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        r = p.add_run(f"[{i}]  {ref}")
        _set_run_font(r, size=11)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build()
    print(out)
